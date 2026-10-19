"""Superintendencia Financiera de Colombia - Circular Basica Juridica (CBJ): attachments of every Capitulo.

    Parte  ->  Titulo  ->  Capitulo  ->  attachments (PDF, DOC/DOCX, XLS/XLSX, ...)

The page is a shell around an iframe (".../visualizadorCBJ/...") that runs a single-page app. Everything is loaded on demand
by the app's own API calls (they need a session key the app obtains itself), so the crawl drives the real page with
Playwright, exactly as a person would:

* the sidebar (a table of Parte / Titulo / Capitulo rows) is clicked - never the "DESCARGAR" buttons;
* nothing is waited for with fixed sleeps: the crawl waits for *conditions* (the clicked Capitulo is the selected row, the
  app's own requests have finished and stayed quiet, a sidebar row has changed ...);
* attachments are collected from three places and cross-checked: the app's attachment API responses (keyed by the Parte /
  Titulo / Capitulo numbers they carry), the DOM of the main page and of every iframe, and the real request the page makes when
  an attachment is clicked (that request is recorded and cancelled, so nothing is downloaded);
* a Capitulo that fails is retried, logged and skipped; the crawl continues.

Run:  python scripts/cbj_attachments.py --out output/cbj_attachments.xlsx
"""
from __future__ import annotations

import argparse
import json
import logging
import re
import sys
import time
from dataclasses import asdict, dataclass, field
from pathlib import Path
from urllib.parse import urljoin, urlparse

START_URL = "https://www.superfinanciera.gov.co/publicaciones/10115528/circular-basica-juridica-ce-00625/"
log = logging.getLogger("cbj")

DOC_EXT = {"pdf", "doc", "docx", "xls", "xlsx", "xlsm", "csv", "ppt", "pptx", "zip", "rar", "txt", "rtf", "odt", "ods", "xml", "json"}
MIME_EXT = {"application/pdf": "pdf", "application/msword": "doc", "application/vnd.ms-excel": "xls",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
            "application/vnd.ms-powerpoint": "ppt", "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
            "application/zip": "zip", "text/csv": "csv", "text/plain": "txt"}
LEVELS = {"3": "Capitulo", "4": "Numeral", "5": "Subnumeral"}


@dataclass
class Site:
    """Everything particular to this site - change here if the page changes."""
    viewer_iframe: str = "iframe[src*='visualizadorCBJ']"
    start_card_arrow: str = "button.flecha"                      # the arrow that opens a Parte (NOT the DESCARGAR button)
    sidebar_row: str = "#fixedColumn tr"
    capitulo_row: str = "#fixedColumn tr#capitulos"
    api_pattern: str = r"/sfcservices/"                          # the site's own requests: "the page is still loading" while any is open
    anexos_api: str = "consultarAnexosElemento"
    titulo_api: str = "consultarTituloFecha"                     # requested whenever a Titulo is opened: .../consultarTituloFecha/<titulo>/<parte>/false
    # NB the site writes its numbers deepest first: consultarCapituloFecha/<capitulo>/<titulo>/<parte>, consultarTituloFecha/<titulo>/<parte>
    attachment_item: str = "#anexos"                             # clickable attachment in a Capitulo's content
    attachment_name: str = ".anexo"
    download_pattern: str = r"/descargarAnexo/"                  # the request an attachment click makes
    forbidden_text: str = "descargar"                            # nothing whose text contains this is ever clicked


@dataclass
class Timing:
    """Upper bounds for the wait conditions, in seconds. A condition ends its wait as soon as it is true."""
    page: float = 60
    viewer: float = 45
    rows: float = 20
    chapter: float = 45
    click_capture: float = 8
    idle_quiet_ms: int = 600
    retries: int = 2


@dataclass
class Attachment:
    parte: str
    titulo: str
    capitulo: str
    name: str
    url: str
    ext: str = ""
    content_type: str = ""
    attachment_id: str = ""
    level: str = "Capitulo"                    # which part of the Capitulo the attachment belongs to
    numeral: str = ""
    subnumeral: str = ""
    found_in: str = ""                         # api / dom / iframe-dom / click
    verified_by_click: bool = False
    page_url: str = ""
    title_row: int = 1                         # the site lists a Titulo twice under one Parte: 1 = first row, 2 = second row ...
    app_ids: str = ""                          # parte/titulo/capitulo numbers the site's own API uses for this Capitulo
    size_bytes: int | None = None              # filled only when the files are read (read_files)
    pages: str = ""
    language: str = ""
    last_modified: str = ""
    read_error: str = ""

    @property
    def hierarchy(self) -> str:
        titulo = self.titulo + (f" [title row {self.title_row}]" if self.title_row > 1 else "")
        return " > ".join(x for x in (self.parte, titulo, self.capitulo) if x)


@dataclass
class ChapterRun:
    parte: str
    titulo: str
    capitulo: str
    title_row: int = 1
    app_ids: str = ""
    check: str = ""                            # cross-check of the site's attachment data against the clicks
    status: str = "pending"                    # ok / failed / duplicate (a repeated sidebar row with the same site ids: not read twice)
    attachments: int = 0
    error: str = ""
    seconds: float = 0.0
    attempts: int = 0


# ── pure helpers (no browser) ────────────────────────────────────────────────
_ROMAN = {"I": 1, "V": 5, "X": 10, "L": 50, "C": 100, "D": 500, "M": 1000}


def roman_to_int(s: str) -> int:
    s = s.strip().upper()
    if s.isdigit():
        return int(s)
    total = 0
    for i, ch in enumerate(s):
        v = _ROMAN.get(ch)
        if v is None:
            raise ValueError(f"not a roman numeral: {s!r}")
        total += -v if i + 1 < len(s) and _ROMAN.get(s[i + 1], 0) > v else v
    return total


def ordinal_of(label: str) -> int | None:
    """'Capitulo III. Gobierno corporativo' -> 3 ; 'Parte II. ...' -> 2 ; None when the label has no number."""
    m = re.match(r"^\s*(?:parte|t[ií]tulo|cap[ií]tulo)\s+([IVXLCDM]+|\d+)\b", label, re.I)
    try:
        return roman_to_int(m.group(1)) if m else None
    except ValueError:
        return None


def parse_sidebar(rows: list[dict]) -> list[dict]:
    """Sidebar rows -> the hierarchy of every Capitulo row: [{parte, titulo, titulo_occ, capitulo, row}].

    A row is a Parte when its button is 'flechaParte', a Titulo when its id is 'titulos', a Capitulo when its id is 'capitulos'."""
    out, parte, titulo, occ, seen = [], "", "", 0, {}
    for r in rows:
        text = (r.get("text") or "").strip()
        if r.get("btn") == "flechaParte":
            parte, titulo, occ, seen = text, "", 0, {}
        elif r.get("id") == "titulos":
            titulo = text
            occ = seen[text] = seen.get(text, -1) + 1             # the site can list the same Titulo label twice under one Parte
        elif r.get("id") == "capitulos" and text:
            out.append({"parte": parte, "titulo": titulo, "titulo_occ": occ, "capitulo": text, "row": r.get("i")})
    return out


def file_ext(name: str = "", url: str = "", content_type: str = "") -> str:
    """Extension from the file name, else the URL path, else the content type."""
    for src in (name, urlparse(url).path):
        m = re.search(r"\.([A-Za-z0-9]{2,5})\s*$", src or "")
        if m and m.group(1).lower() in DOC_EXT:
            return m.group(1).lower()
    return MIME_EXT.get((content_type or "").split(";")[0].strip().lower(), "")


def absolute(base: str, href: str) -> str:
    return urljoin(base, (href or "").strip()) if href else ""


def dedupe(items: list[Attachment]) -> list[Attachment]:
    """One row per (Parte, Titulo, Capitulo, URL). The same file listed twice in one Capitulo is kept once (what was learned
    from each sighting is merged); the same file in two Capitulos keeps one row for each - the hierarchy is never lost."""
    seen: dict[tuple, Attachment] = {}
    for a in items:
        key = (a.parte, a.titulo, a.title_row, a.capitulo, a.url or a.name)
        cur = seen.get(key)
        if cur is None:
            seen[key] = a
            continue
        cur.found_in = ", ".join(dict.fromkeys(x for x in (cur.found_in + ", " + a.found_in).split(", ") if x))
        cur.verified_by_click = cur.verified_by_click or a.verified_by_click
        for f in ("content_type", "attachment_id", "numeral", "subnumeral", "ext", "name", "pages", "language", "last_modified", "read_error"):
            if not getattr(cur, f) and getattr(a, f):
                setattr(cur, f, getattr(a, f))
        if cur.size_bytes is None and a.size_bytes is not None:
            cur.size_bytes = a.size_bytes
        if cur.level == "Capitulo" and a.level != "Capitulo":
            cur.level = a.level
    return list(seen.values())


def unique_urls(items: list[Attachment]) -> list[dict]:
    """Each URL once, with every place of the hierarchy it appears in."""
    by_url: dict[str, dict] = {}
    for a in items:
        d = by_url.setdefault(a.url or a.name, {"url": a.url, "name": a.name, "ext": a.ext, "content_type": a.content_type, "locations": []})
        if a.hierarchy not in d["locations"]:
            d["locations"].append(a.hierarchy)
    return [{**d, "occurrences": len(d["locations"]), "locations": " | ".join(d["locations"])} for d in by_url.values()]


# ── in-page scripts ──────────────────────────────────────────────────────────
def tracker_script(pattern: str) -> str:
    """Runs in every frame before the page's own scripts and counts the site's own XHR/fetch calls that are in flight, so
    'the page has finished loading' is a condition that can be waited for. (Long-lived third-party connections such as
    Firestore's are not counted: they never end.)"""
    return """(() => {
  if (window.__net) return;
  const st = window.__net = { inflight: 0, last: Date.now(), re: new RegExp(%s) };
  const done = () => { st.inflight = Math.max(0, st.inflight - 1); st.last = Date.now(); };
  const open = XMLHttpRequest.prototype.open, send = XMLHttpRequest.prototype.send;
  XMLHttpRequest.prototype.open = function (m, u) { this.__u = String(u); return open.apply(this, arguments); };
  XMLHttpRequest.prototype.send = function () {
    if (st.re.test(this.__u || '')) { st.inflight++; st.last = Date.now(); this.addEventListener('loadend', done, { once: true }); }
    return send.apply(this, arguments); };
  const f = window.fetch;
  if (f) window.fetch = function (input) {
    const u = typeof input === 'string' ? input : (input && input.url) || '';
    if (!st.re.test(u)) return f.apply(this, arguments);
    st.inflight++; st.last = Date.now();
    return f.apply(this, arguments).finally(done); };
})();""" % json.dumps(pattern)


JS_IDLE = "(q) => !!window.__net && window.__net.inflight === 0 && (Date.now() - window.__net.last) >= q"
JS_HAS_TRACKER = "() => !!window.__net"
JS_SIDEBAR = """(sel) => [...document.querySelectorAll(sel)].map((tr, i) => {
  const td = tr.querySelector('td'); const btn = tr.querySelector('button');
  return { i, id: tr.id, selected: tr.classList.contains('selected'), text: (td ? td.innerText : '').replace(/\\s+/g, ' ').trim(),
           btn: btn ? btn.id : null, rotated: btn ? btn.classList.contains('rotate-180') : null }; })"""
JS_SIGNATURE = """(sel) => JSON.stringify([...document.querySelectorAll(sel)].map(tr => { const td = tr.querySelector('td'), b = tr.querySelector('button');
  return [(td ? td.innerText : '').replace(/\\s+/g, ' ').trim(), b ? b.classList.contains('rotate-180') : null, tr.classList.contains('selected')]; }))"""
JS_SIGNATURE_CHANGED = "([sel, before]) => (" + JS_SIGNATURE + ")(sel) !== before"
JS_CHAPTER_SELECTED = """([sel, cap]) => { const t = document.querySelector(sel + '.selected'); return !!t && (t.innerText || '').replace(/\\s+/g, ' ').trim() === cap; }"""
JS_DOM_DOCS = """(exts) => {
  const out = [], t = e => (e.innerText || e.textContent || '').replace(/\\s+/g, ' ').trim();
  const re = new RegExp('\\\\.(' + exts.join('|') + ')(?:$|[?#&])', 'i');
  for (const a of document.querySelectorAll('a[href], area[href]')) {
    const h = a.getAttribute('href') || '';
    if (!/^(javascript:|#|mailto:|tel:)/i.test(h) && (re.test(h) || a.hasAttribute('download'))) out.push({ name: t(a) || a.getAttribute('download') || '', href: h });
  }
  for (const e of document.querySelectorAll('iframe[src], embed[src], object[data], source[src]')) {
    const h = e.getAttribute('src') || e.getAttribute('data') || '';
    if (re.test(h)) out.push({ name: '', href: h });
  }
  return out; }"""


# ── the crawl ────────────────────────────────────────────────────────────────
class Cancelled(Exception):
    """The user stopped the crawl."""


class CbjCrawler:
    def __init__(self, page, site: Site | None = None, timing: Timing | None = None, click_verify: bool = True, read_files: bool = False,
                 progress=None, cancelled=None):
        self.page, self.site, self.t, self.click_verify = page, site or Site(), timing or Timing(), click_verify
        self.read_files = read_files                             # download each attachment INSIDE the session (into memory) to learn size / pages / date
        self.progress, self.cancelled = progress, cancelled
        self.frame = None
        self.start_url = ""
        self.anexos: dict[tuple, dict[str, dict]] = {}          # (parte, titulo, capitulo) -> {id: record}, from the app's own API responses
        self.api_root = ""                                       # ".../apiCircularBasica/", read from a real response URL
        self.cap_calls: list[tuple] = []                         # (parte, titulo, capitulo) the app loaded, in order (from consultarCapituloFecha)
        self.title_keys: dict[tuple, tuple] = {}                 # (parte label, titulo label, row #) -> the site's own (parte, titulo) numbers
        self.last_check = ""
        self.cancel_downloads = False

    # -- wiring ---------------------------------------------------------------
    def attach(self, context) -> None:
        """Listeners and the wait-condition tracker. Call once, before the first navigation."""
        context.add_init_script(tracker_script(self.site.api_pattern))
        self.page.on("response", self._on_response)
        context.route(re.compile(self.site.download_pattern), self._on_download_request)

    def _on_response(self, resp) -> None:
        m = re.search(r"consultarCapituloFecha/(\d+)/(\d+)/(\d+)", resp.url)
        if m:
            self.cap_calls.append((m.group(3), m.group(2), m.group(1)))       # written capitulo/titulo/parte; kept as (parte, titulo, capitulo)
            return
        if self.site.anexos_api not in resp.url:
            return
        self.api_root = self.api_root or resp.url.split(self.site.anexos_api)[0]
        try:
            data = resp.json()
        except Exception:
            return
        q = dict(re.findall(r"[?&](\w+)=([^&]*)", resp.url))
        key = (q.get("parte", ""), q.get("titulo", ""), q.get("capitulo", ""))
        for rec in data if isinstance(data, list) else []:
            if isinstance(rec, dict) and rec.get("id") is not None:
                self.anexos.setdefault(key, {})[str(rec["id"])] = {**rec, "_element": q.get("element", "")}

    def _on_download_request(self, route) -> None:
        """While an attachment is being clicked for verification the request is cancelled: its address is recorded, nothing is downloaded."""
        if self.cancel_downloads:
            return route.abort()
        route.continue_()

    # -- waits (conditions, not sleeps) ----------------------------------------
    def idle(self, timeout: float | None = None) -> bool:
        """The site's own requests (in every frame that has the tracker) have finished and stayed quiet."""
        timeout = timeout or self.t.chapter
        deadline = time.monotonic() + timeout
        ok = True
        for fr in list(self.page.frames):
            if fr.is_detached() or not fr.url.startswith("http"):
                continue
            try:
                if not fr.evaluate(JS_HAS_TRACKER):
                    continue
                fr.wait_for_function(JS_IDLE, arg=self.t.idle_quiet_ms, timeout=max(0.5, deadline - time.monotonic()) * 1000)
            except Exception:
                ok = False
        if not ok:
            log.warning("the page was still busy after %.0fs; continuing", timeout)
        return ok

    def say(self, **kw) -> None:
        if self.progress:
            self.progress(kw)

    def check_cancelled(self) -> None:
        if self.cancelled and self.cancelled():
            raise Cancelled()

    def rows(self) -> list[dict]:
        return self.frame.evaluate(JS_SIDEBAR, self.site.sidebar_row)

    def _signature(self) -> str:
        return self.frame.evaluate(JS_SIGNATURE, self.site.sidebar_row)

    def _wait_sidebar_change(self, before: str) -> None:
        try:
            self.frame.wait_for_function(JS_SIGNATURE_CHANGED, arg=[self.site.sidebar_row, before], timeout=self.t.rows * 1000)
        except Exception:
            pass                                                 # nothing changed (e.g. a Titulo without Capitulos); the idle wait that follows decides

    # -- opening the viewer -----------------------------------------------------
    def open(self, url: str) -> None:
        self.start_url = url
        self.page.goto(url, wait_until="domcontentloaded", timeout=self.t.page * 1000)
        handle = self.page.wait_for_selector(self.site.viewer_iframe, state="attached", timeout=self.t.viewer * 1000)
        self.frame = handle.content_frame()
        self.frame.wait_for_selector(self.site.start_card_arrow, state="attached", timeout=self.t.viewer * 1000)
        self.idle(self.t.viewer)

    def enter_reader(self) -> None:
        """Open the first Parte from the start screen with its arrow (never the DESCARGAR button); the sidebar then lists every Parte."""
        arrow = self.frame.locator(self.site.start_card_arrow).first
        if self.site.forbidden_text in (arrow.inner_text() or "").lower():
            raise RuntimeError("refusing to click a DESCARGAR button")
        arrow.dispatch_event("click")
        self.frame.wait_for_selector(self.site.sidebar_row, state="attached", timeout=self.t.viewer * 1000)
        self.idle(self.t.rows)

    def recover(self) -> None:
        """Start again from the page address (after a failure): the sidebar is rebuilt from scratch."""
        log.info("re-opening the page")
        self.open(self.start_url)
        self.enter_reader()

    # -- sidebar navigation -----------------------------------------------------
    def _click_row(self, i: int, button: bool) -> None:
        row = self.frame.locator(self.site.sidebar_row).nth(i)
        target = row.locator("button").first if button else row.locator("td").first
        if self.site.forbidden_text in (target.inner_text() or "").lower():
            raise RuntimeError("refusing to click a DESCARGAR control")
        target.dispatch_event("click")                           # a cookie banner may overlay the frame; the event still reaches the row

    def _open(self, find) -> None:
        for _ in range(3):
            row = find(self.rows())
            if row is None:
                raise LookupError("not found in the sidebar")
            if row["rotated"]:
                return
            before = self._signature()
            self._click_row(row["i"], button=True)
            self._wait_sidebar_change(before)
            self.idle(self.t.rows)

    def open_parte(self, parte: str) -> None:
        try:
            self._open(lambda rows: next((r for r in rows if r["btn"] == "flechaParte" and r["text"] == parte), None))
        except LookupError:
            raise LookupError(f"Parte not found in the sidebar: {parte}") from None

    def open_titulo(self, parte: str, titulo: str, occ: int = 0) -> None:
        """Open a Titulo and learn the site's own (parte, titulo) numbers from the request the page makes for it. (Labels cannot be
        trusted for that: the site lists one Titulo label twice under a Parte, and the two rows are different Titulos.)"""
        ident = (parte, titulo, occ)
        for _ in range(4):
            row = self._titulo_row(self.rows(), parte, titulo, occ)
            if row is None:
                raise LookupError(f"Titulo not found in the sidebar: {parte} > {titulo} (row {occ + 1})")
            if row["rotated"] and ident in self.title_keys:
                return
            before = self._signature()
            if row["rotated"]:                                   # open already (the first Titulo opens by itself) but its numbers are unknown: close, then reopen
                self._click_row(row["i"], button=True)
                self._wait_sidebar_change(before)
                self.idle(self.t.rows)
                continue
            try:
                with self.page.expect_response(lambda r: self.site.titulo_api in r.url, timeout=self.t.rows * 1000) as info:
                    self._click_row(row["i"], button=True)
                m = re.search(self.site.titulo_api + r"/(\d+)/(\d+)", info.value.url)
                if m:
                    self.title_keys[ident] = (m.group(2), m.group(1))            # written titulo/parte; kept as (parte, titulo)
            except Exception:
                log.debug("no %s request seen after opening %s", self.site.titulo_api, titulo[:40])
            self._wait_sidebar_change(before)
            self.idle(self.t.rows)
        raise LookupError(f"Titulo could not be opened: {parte} > {titulo} (row {occ + 1})")

    @staticmethod
    def _titulo_row(rows: list[dict], parte: str, titulo: str, occ: int = 0) -> dict | None:
        """The occ-th (0-based) Titulo row with this label under this Parte."""
        cur, n = "", -1
        for r in rows:
            if r["btn"] == "flechaParte":
                cur, n = r["text"], -1
            elif r["id"] == "titulos" and cur == parte and r["text"] == titulo:
                n += 1
                if n == occ:
                    return r
        return None

    def structure(self, only_parte: int | None = None) -> list[dict]:
        """Open every Parte (or just one) and its Titulos in turn and return the hierarchy of the Capitulos."""
        found: list[dict] = []
        for parte in [r["text"] for r in self.rows() if r["btn"] == "flechaParte"]:
            if only_parte and ordinal_of(parte) != only_parte:
                continue
            self.check_cancelled()
            self.say(phase="crawl", message=f"Reading the menu: {parte[:60]}", links=0)
            self.open_parte(parte)
            titulos, cur, seen = [], "", {}
            for r in self.rows():
                if r["btn"] == "flechaParte":
                    cur = r["text"]
                elif r["id"] == "titulos" and cur == parte:
                    occ = seen[r["text"]] = seen.get(r["text"], -1) + 1
                    titulos.append((r["text"], occ))
            for titulo, occ in titulos:
                self.check_cancelled()
                self.open_titulo(parte, titulo, occ)
                caps = [c for c in parse_sidebar(self.rows()) if (c["parte"], c["titulo"], c["titulo_occ"]) == (parte, titulo, occ)]
                found += caps
                log.info("%s > %s%s: %d capitulo(s)", parte[:24], titulo[:44], f" (title row {occ + 1})" if occ else "", len(caps))
        return found

    # -- one Capitulo -------------------------------------------------------------
    def read_chapter(self, parte: str, titulo: str, capitulo: str, occ: int = 0) -> list[Attachment]:
        self.open_parte(parte)
        self.open_titulo(parte, titulo, occ)
        row = next((c for c in parse_sidebar(self.rows()) if (c["parte"], c["titulo"], c["titulo_occ"], c["capitulo"]) == (parte, titulo, occ, capitulo)), None)
        if row is None:
            raise LookupError(f"Capitulo not found in the sidebar: {capitulo}")
        selected = next((r for r in self.rows() if r["i"] == row["row"]), {}).get("selected")
        mark = len(self.cap_calls)
        if not selected:
            self._click_row(row["row"], button=False)
        # the condition that matters: the clicked Capitulo is the selected row; then the page's own requests settle
        self.frame.wait_for_function(JS_CHAPTER_SELECTED, arg=[self.site.capitulo_row, capitulo], timeout=self.t.chapter * 1000)
        self.idle()
        # the site's own numbers: (parte, titulo) learnt when the Titulo was opened + this Capitulo's number
        pt = self.title_keys.get((parte, titulo, occ)) or (str(ordinal_of(parte) or ""), str(ordinal_of(titulo) or ""))
        key = (pt[0], pt[1], str(ordinal_of(capitulo) or ""))
        if key not in self.cap_calls[mark:] and key not in self.cap_calls:
            log.warning("%s: the page did not load the Capitulo %s that was expected (loaded: %s)", capitulo[:40], "/".join(key), self.cap_calls[mark:][-1:] or "nothing")
        api_rows = self._from_api(parte, titulo, capitulo, key, occ + 1)
        dom_rows = self._from_dom(parte, titulo, capitulo, occ + 1, key)
        click_rows = self._verify_by_click(parte, titulo, capitulo, occ + 1, key) if self.click_verify else []
        self.last_check = self._cross_check(api_rows, click_rows) if self.click_verify else "api only (clicks disabled)"
        return dedupe(api_rows + dom_rows + click_rows)

    @staticmethod
    def _cross_check(api_rows: list[Attachment], click_rows: list[Attachment]) -> str:
        """The site's attachment data against what is displayed and clickable: both lists should name the same files."""
        a = {r.attachment_id for r in api_rows if r.attachment_id}
        c = {r.attachment_id for r in click_rows if r.attachment_id}
        unclickable = sum(1 for r in click_rows if not r.attachment_id)
        if a == c and not unclickable:
            return f"match ({len(a)})" if a else "none"
        parts = []
        if a - c:
            parts.append(f"in the site's data but not on the page: {len(a - c)}")
        if c - a:
            parts.append(f"on the page but not in the site's data: {len(c - a)}")
        if unclickable:
            parts.append(f"shown but without a download request: {unclickable}")
        return "; ".join(parts)

    def _from_api(self, parte: str, titulo: str, capitulo: str, key: tuple, title_row: int) -> list[Attachment]:
        """The app's own attachment records for this Capitulo (and its numerals / subnumerals), matched by the numbers in the request."""
        out = []
        for rid, rec in self.anexos.get(tuple(key), {}).items():
            name = str(rec.get("nombre") or "")
            numeral, sub = str(rec.get("numeral") or ""), str(rec.get("subnumeral") or "")
            out.append(Attachment(parte, titulo, capitulo, name, urljoin(self.api_root, f"descargarAnexo/{rid}"), file_ext(name, "", rec.get("contentType", "")),
                                  str(rec.get("contentType") or ""), rid, LEVELS.get(str(rec.get("_element") or rec.get("elemento") or ""), "Capitulo"),
                                  "" if numeral in ("", "0") else numeral, "" if sub in ("", "0") else sub, "api", False, self.frame.url,
                                  title_row, "/".join(key)))
        return out

    def _from_dom(self, parte: str, titulo: str, capitulo: str, title_row: int, key: tuple) -> list[Attachment]:
        """Document links and framed documents in the main page AND in every iframe; relative addresses become absolute."""
        out = []
        for fr in self.page.frames:
            if fr.is_detached() or not fr.url.startswith("http"):
                continue
            try:
                docs = fr.evaluate(JS_DOM_DOCS, sorted(DOC_EXT))
            except Exception:
                continue
            for d in docs:
                url = absolute(fr.url, d["href"])
                name = d["name"] or url.rsplit("/", 1)[-1].split("?")[0]
                out.append(Attachment(parte, titulo, capitulo, name, url, file_ext(name, url), "", "", "Capitulo", "", "",
                                      "dom" if fr == self.page.main_frame else "iframe-dom", False, fr.url, title_row, "/".join(key)))
        return out

    def _verify_by_click(self, parte: str, titulo: str, capitulo: str, title_row: int, key: tuple) -> list[Attachment]:
        """Click every attachment shown in the content and record the request the page really makes.

        By default that request is cancelled (nothing is downloaded). With read_files the request is allowed and the file is read
        into memory - exactly what clicking it in a browser does - to learn its size, page count and date; it is never saved."""
        out = []
        items = self.frame.locator(self.site.attachment_item)
        pattern = re.compile(self.site.download_pattern)
        for n in range(items.count()):
            item = items.nth(n)
            try:
                name = (item.locator(self.site.attachment_name).first.inner_text(timeout=2000) or "").strip()
            except Exception:
                name = (item.inner_text() or "").strip()
            if self.site.forbidden_text in name.lower():
                continue
            self.cancel_downloads = not self.read_files
            resp = None
            try:
                if self.read_files:
                    with self.page.expect_response(lambda r: bool(pattern.search(r.url)), timeout=self.t.click_capture * 1000) as info:
                        item.dispatch_event("click")
                    resp = info.value
                    url = resp.url
                else:
                    # 'requestfailed' is only fired once our route has cancelled the request, so the flag cannot be reset too early
                    with self.page.expect_event("requestfailed", predicate=lambda r: bool(pattern.search(r.url)), timeout=self.t.click_capture * 1000) as info:
                        item.dispatch_event("click")
                    url = info.value.url
            except Exception:
                log.warning("clicking %r produced no download request", name)
                out.append(Attachment(parte, titulo, capitulo, name, "", file_ext(name), "", "", "Capitulo", "", "", "dom", False, self.frame.url, title_row, "/".join(key)))
                continue
            finally:
                self.cancel_downloads = False
            m = re.search(r"/descargarAnexo/(\d+)", url)
            att = Attachment(parte, titulo, capitulo, name, url, file_ext(name, url), "", m.group(1) if m else "", "Capitulo", "", "", "click", True,
                             self.frame.url, title_row, "/".join(key))
            if resp is not None:
                self._read_response(att, resp)
            out.append(att)
        return out

    @staticmethod
    def _read_response(att: Attachment, resp) -> None:
        """Size, page count, language and date of a file the page just downloaded (read from memory; nothing is saved)."""
        from ..analysis.document_analyzer import inspect_document, parse_http_date
        try:
            if resp.status >= 400:
                att.read_error = f"HTTP {resp.status}"
                return
            data = resp.body()
            headers = {k.lower(): v for k, v in resp.headers.items()}
            att.size_bytes = len(data)
            att.content_type = att.content_type or headers.get("content-type", "").split(";")[0]
            att.last_modified = parse_http_date(headers.get("last-modified", "")) or ""
            kind = {"pdf": "PDF", "docx": "DOCX"}.get(att.ext, "")
            if kind:
                info = inspect_document(kind, data)
                att.pages = str(info.pages) if info.pages is not None else "Unknown"
                att.language = info.language if info.language not in ("N/A", None) else ""
                att.last_modified = att.last_modified or (info.modified or "")
                if info.problem:
                    att.read_error = info.problem
            else:
                att.pages = "Unknown"                              # .doc / .xls ... cannot be opened for a page count
        except Exception as e:
            att.read_error = f"{type(e).__name__}: {str(e)[:120]}"


# ── running it ───────────────────────────────────────────────────────────────
@dataclass
class Result:
    attachments: list[Attachment] = field(default_factory=list)
    chapters: list[ChapterRun] = field(default_factory=list)
    started: str = ""
    seconds: float = 0.0
    notes: list[str] = field(default_factory=list)

    @property
    def errors(self) -> list[ChapterRun]:
        return [c for c in self.chapters if c.status == "failed"]


def crawl(url: str = START_URL, *, headless: bool = True, only_parte: int | None = None, max_chapters: int | None = None,
          click_verify: bool = True, timing: Timing | None = None, site: Site | None = None, browser_channel: str | None = None,
          read_files: bool = False, progress=None, cancelled=None) -> Result:
    """Walk Parte -> Titulo -> Capitulo and collect every attachment. A Capitulo that fails is retried, logged and skipped."""
    from playwright.sync_api import sync_playwright
    timing = timing or Timing()
    res = Result(started=time.strftime("%Y-%m-%d %H:%M:%S"))
    t0 = time.monotonic()
    with sync_playwright() as p:
        browser = None
        for ch in ([browser_channel] if browser_channel else ["msedge", "chrome", None]):
            try:
                browser = p.chromium.launch(headless=headless, **({"channel": ch} if ch else {}))
                log.info("browser: %s", ch or "chromium")
                break
            except Exception as e:
                log.debug("browser %s unavailable: %s", ch, str(e).splitlines()[0])
        if browser is None:
            raise RuntimeError("No browser could be started (install Microsoft Edge or Chrome, or run `playwright install chromium`).")
        try:
            ctx = browser.new_context(viewport={"width": 1500, "height": 1000}, accept_downloads=False)
            page = ctx.new_page()
            crawler = CbjCrawler(page, site, timing, click_verify or read_files, read_files, progress, cancelled)
            crawler.attach(ctx)
            crawler.say(phase="crawl", message="Opening the page", links=0)
            crawler.open(url)
            crawler.enter_reader()
            try:
                plan = crawler.structure(only_parte)
            except Cancelled:
                res.notes.append("Stopped by the user.")
                plan = []
            if only_parte:
                plan = [c for c in plan if ordinal_of(c["parte"]) == only_parte]
            if max_chapters:
                plan = plan[:max_chapters]
            log.info("%d capitulo(s) to read", len(plan))
            first_seen: dict[str, str] = {}                           # site ids -> where this Capitulo was first listed
            for n, c in enumerate(plan, 1):
                if cancelled and cancelled():
                    res.notes.append("Stopped by the user.")
                    break
                crawler.say(phase="crawl", message=f"Capítulo {n} of {len(plan)}: {c['capitulo'][:60]}", links=len(res.attachments))
                run = ChapterRun(c["parte"], c["titulo"], c["capitulo"], c.get("titulo_occ", 0) + 1)
                res.chapters.append(run)
                started = time.monotonic()
                pt = crawler.title_keys.get((c["parte"], c["titulo"], c.get("titulo_occ", 0)))
                num = ordinal_of(c["capitulo"])
                run.app_ids = "/".join(pt + (str(num),)) if pt and num else ""
                if run.app_ids and run.app_ids in first_seen:     # the site lists this row twice: same data, so read it once
                    run.status, run.check = "duplicate", f"same site ids as {first_seen[run.app_ids]}; not read again"
                    log.info("[%d/%d] %s > %s > %s -> duplicate row (%s)", n, len(plan), c["parte"][:8], c["titulo"][:20], c["capitulo"][:40], run.check)
                    continue
                if run.app_ids:
                    first_seen[run.app_ids] = f"{c['titulo'][:50]} > {c['capitulo'][:40]}"
                for attempt in range(1, timing.retries + 2):
                    run.attempts = attempt
                    try:
                        atts = crawler.read_chapter(c["parte"], c["titulo"], c["capitulo"], c.get("titulo_occ", 0))
                        run.status, run.attachments, run.error = "ok", len(atts), ""
                        run.check = crawler.last_check
                        if run.check not in ("none",) and not run.check.startswith(("match", "api only")):
                            log.warning("%s: cross-check differs - %s", c["capitulo"][:40], run.check)
                        res.attachments += atts
                        break
                    except Cancelled:
                        run.status = "stopped"
                        break
                    except Exception as e:
                        run.status, run.error = "failed", f"{type(e).__name__}: {str(e).splitlines()[0][:200]}"
                        log.error("[%d/%d] %s | attempt %d failed: %s", n, len(plan), c["capitulo"][:50], attempt, run.error)
                        if attempt <= timing.retries:
                            try:
                                crawler.recover()
                            except Exception as e2:
                                log.error("could not re-open the page: %s", e2)
                run.seconds = round(time.monotonic() - started, 1)
                log.info("[%d/%d] %s > %s > %s -> %s (%d attachment(s), %.1fs)", n, len(plan), c["parte"][:8], c["titulo"][:20], c["capitulo"][:40],
                         run.status, run.attachments, run.seconds)
        finally:
            browser.close()
    res.attachments = dedupe(res.attachments)
    res.seconds = round(time.monotonic() - t0, 1)
    return res


# ── output ───────────────────────────────────────────────────────────────────
COLUMNS = [("parte", "Parte"), ("titulo", "Titulo"), ("capitulo", "Capitulo"), ("level", "Level"), ("numeral", "Numeral"), ("subnumeral", "Subnumeral"),
           ("name", "Attachment"), ("ext", "Type"), ("content_type", "Content-Type"), ("url", "Attachment URL"), ("attachment_id", "Attachment ID"),
           ("found_in", "Found In"), ("verified_by_click", "Verified By Click"), ("title_row", "Title Row #"),
           ("app_ids", "Site IDs (Parte/Titulo/Capitulo)"), ("page_url", "Page / Frame URL")]


def _safe(v):
    return "'" + v if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") else v


READ_COLUMNS = [("size_bytes", "Size (bytes)"), ("pages", "Pages"), ("language", "Language"), ("last_modified", "Last Modified"), ("read_error", "Read Error")]


def write_outputs(res: Result, out: Path) -> list[Path]:
    """Excel workbook (Attachments, Unique URLs, Capitulos, Errors, Summary), plus a JSON copy."""
    import pandas as pd
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    out.parent.mkdir(parents=True, exist_ok=True)
    cols = COLUMNS + (READ_COLUMNS if any(a.size_bytes is not None or a.read_error for a in res.attachments) else [])
    att = pd.DataFrame([{label: _safe(getattr(a, key)) for key, label in cols} for a in res.attachments], columns=[c[1] for c in cols])
    uniq = pd.DataFrame(unique_urls(res.attachments), columns=["name", "ext", "content_type", "url", "occurrences", "locations"]).rename(
        columns={"name": "Attachment", "ext": "Type", "content_type": "Content-Type", "url": "Attachment URL", "occurrences": "Occurrences", "locations": "Found In (hierarchy)"})
    chap = pd.DataFrame([asdict(c) for c in res.chapters]).rename(columns=str.title) if res.chapters else pd.DataFrame()
    errs = chap[chap["Status"] == "failed"] if not chap.empty else chap
    summary = pd.DataFrame([("Started", res.started), ("Duration (s)", res.seconds),
                            ("Capitulos read", sum(1 for c in res.chapters if c.status == "ok")),
                            ("Capitulos listed twice by the site (not read again)", sum(1 for c in res.chapters if c.status == "duplicate")),
                            ("Capitulos failed", len(res.errors)), ("Attachment rows (one per Capitulo + file)", len(res.attachments)),
                            ("Unique attachment URLs", len(uniq)),
                            ("Note", "Attachment URLs are the site's download endpoints; they only work inside the viewer's own session (a plain request is answered 401).")],
                           columns=["Item", "Value"])
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        for name, df in (("Summary", summary), ("Attachments", att), ("Unique URLs", uniq), ("Capitulos", chap), ("Errors", errs)):
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            for c in ws[1]:
                c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
            for i, col in enumerate(df.columns, 1):
                width = min(70, max(10, int(df[col].astype(str).str.len().quantile(0.9)) + 2 if len(df) else 12, len(str(col)) + 2))
                ws.column_dimensions[get_column_letter(i)].width = width
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
    js = out.with_suffix(".json")
    js.write_text(json.dumps({"summary": dict(summary.values.tolist()), "attachments": [asdict(a) | {"hierarchy": a.hierarchy} for a in res.attachments],
                              "chapters": [asdict(c) for c in res.chapters]}, ensure_ascii=False, indent=1), encoding="utf-8")
    return [out, js]


def setup_logging(log_file: Path | None, verbose: bool = False) -> None:
    fmt = logging.Formatter("%(asctime)s %(levelname)-7s %(message)s", "%H:%M:%S")
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")      # accents in the console on Windows
        except Exception:
            pass
    root = logging.getLogger("cbj")
    root.setLevel(logging.DEBUG if verbose else logging.INFO)
    root.handlers.clear()
    for h in [logging.StreamHandler(sys.stdout)] + ([logging.FileHandler(log_file, encoding="utf-8")] if log_file else []):
        h.setFormatter(fmt)
        root.addHandler(h)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Attachments of every Capitulo of the Circular Basica Juridica (Superfinanciera).")
    ap.add_argument("--url", default=START_URL)
    ap.add_argument("--out", default=f"output/cbj_attachments_{time.strftime('%Y%m%d')}.xlsx", help="Excel workbook to write (a .json copy is written next to it)")
    ap.add_argument("--log", default="logs/cbj_attachments.log")
    ap.add_argument("--headed", action="store_true", help="show the browser")
    ap.add_argument("--parte", type=int, help="only this Parte (1, 2, 3)")
    ap.add_argument("--max-capitulos", type=int, help="stop after this many Capitulos (for a quick test)")
    ap.add_argument("--no-click-verify", action="store_true", help="do not click attachments to confirm their address (faster)")
    ap.add_argument("--read-files", action="store_true", help="also read each file inside the browser session (into memory, nothing is saved) for size, pages and date")
    ap.add_argument("--retries", type=int, default=2)
    ap.add_argument("-v", "--verbose", action="store_true")
    a = ap.parse_args(argv)
    Path(a.log).parent.mkdir(parents=True, exist_ok=True)
    setup_logging(Path(a.log), a.verbose)
    res = crawl(a.url, headless=not a.headed, only_parte=a.parte, max_chapters=a.max_capitulos, click_verify=not a.no_click_verify, timing=Timing(retries=a.retries), read_files=a.read_files)
    for p in write_outputs(res, Path(a.out)):
        log.info("wrote %s", p)
    log.info("done: %d capitulo(s) (%d read, %d repeated rows skipped), %d failed, %d attachment row(s), %d unique URL(s) in %.0fs",
             len(res.chapters), sum(1 for c in res.chapters if c.status == "ok"), sum(1 for c in res.chapters if c.status == "duplicate"),
             len(res.errors), len(res.attachments), len(unique_urls(res.attachments)), res.seconds)
    for c in res.errors:
        log.error("FAILED: %s > %s > %s : %s", c.parte, c.titulo, c.capitulo, c.error)
    return 1 if res.errors else 0


if __name__ == "__main__":
    sys.exit(main())
