"""Write the 7-tab Excel workbook."""
from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from ..logger import log
from . import config as C
from .dataset import BacklogData
from .filters import public_view, view_mask
from .loader import BacklogError
from .normalize import format_report_date, is_text
from .pivot import build_pivot

_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def safe_sheet_name(name: str, used: set[str] | None = None) -> str:
    """Valid Excel sheet name: no illegal chars, <= 31 chars, unique (case-insensitive)."""
    base = _BAD_SHEET_CHARS.sub("-", name).strip().strip("'") or "Sheet"
    base = base[:31]
    used = used if used is not None else set()
    candidate, i = base, 2
    while candidate.lower() in used:
        suffix = f" ({i})"
        candidate = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(candidate.lower())
    return candidate


def default_output_path(data: BacklogData, output_dir: str = "output") -> Path:
    return Path(output_dir) / f"{data.output_stem}_Report.xlsx"


def _safe_text(v):
    """Strip illegal characters; text starting with '=' must not become an Excel formula."""
    if not isinstance(v, str):
        return v
    v = _ILLEGAL.sub("", v)
    return "'" + v if v.startswith("=") else v


def _clean(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    seen: dict[str, int] = {}
    cols = []
    for c in df.columns:
        name = _ILLEGAL.sub("", str(c)).strip() or "Column"
        seen[name] = seen.get(name, 0) + 1
        cols.append(name if seen[name] == 1 else f"{name}.{seen[name] - 1}")
    df.columns = cols
    for c in df.columns:
        if is_text(df[c]):
            df[c] = df[c].map(_safe_text)
    return df


def _write_table_sheet(writer, sheet: str, df: pd.DataFrame, table_name: str) -> None:
    df = _clean(df)
    df.to_excel(writer, sheet_name=sheet, index=False)
    ws = writer.sheets[sheet]
    ws.freeze_panes = "A2"
    for i, col in enumerate(df.columns, 1):
        sample = df[col].head(200).fillna("").astype(str).str.len().max() if len(df) else 0
        width = max(len(str(col)) + 2, min(int(sample if pd.notna(sample) else 0) + 2, 45), 10)
        ws.column_dimensions[get_column_letter(i)].width = min(width, 45)
        ws.cell(1, i).alignment = Alignment(vertical="center", horizontal="left")
    if len(df):
        tbl = Table(displayName=table_name, ref=f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}")
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tbl)  # tables carry the filter dropdowns


def write_report(data: BacklogData, output_path: str | Path) -> Path:
    out = Path(output_path)
    used: set[str] = set()
    try:
        out.parent.mkdir(parents=True, exist_ok=True)
        with pd.ExcelWriter(out, engine="openpyxl", datetime_format="yyyy-mm-dd hh:mm",
                            date_format="yyyy-mm-dd") as writer:
            ui_sheet = safe_sheet_name(f"UI {format_report_date(data.report_date)}", used)
            _write_table_sheet(writer, ui_sheet, data.ui, "UIData")
            export_sheet = safe_sheet_name(C.SHEET_EXPORT, used)
            _write_table_sheet(writer, export_sheet, public_view(data.export), "ExportData")
            pivot_ws = writer.book.create_sheet(safe_sheet_name(C.SHEET_PIVOT, used))
            build_pivot(pivot_ws, data.export)
            for i, (name, view) in enumerate(C.ACTION_SHEETS.items(), start=4):
                subset = public_view(data.export[view_mask(data.export, view)])
                _write_table_sheet(writer, safe_sheet_name(name, used), subset, f"Detail{i}")
    except PermissionError as e:
        raise BacklogError(f"Could not write '{out.name}'. If it is open in Excel, close it and try again.") from e
    except OSError as e:
        raise BacklogError(f"Could not write the Excel report: {e}") from e
    log.info("Workbook generated: %s", out)
    return out
