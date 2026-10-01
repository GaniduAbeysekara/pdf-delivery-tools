"""Pivot sheet: eight compact summary tables laid out horizontally."""
from __future__ import annotations

import pandas as pd
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config as C

YELLOW = PatternFill("solid", fgColor="FFFF00")
PINK = PatternFill("solid", fgColor="FFD9E4")
RED = PatternFill("solid", fgColor="FF0000")
HEADER_FILL = PatternFill("solid", fgColor="4A6FA5")
THIN = Side(style="thin", color="BBBBBB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CATEGORY_STYLE = {
    C.CAT_CHANGE: (YELLOW, "000000"), C.CAT_POTENTIAL: (YELLOW, "000000"),
    C.CAT_EVC_OPEN: (PINK, "000000"), C.CAT_EVC: (PINK, "000000"),
    C.CAT_DEV: (RED, "FFFFFF"),
}

Row = tuple[str, int, str]  # label, count, category


def _by_api(df: pd.DataFrame) -> list[Row]:
    if df.empty:
        return []
    g = df.groupby(["_api_label", "_cat"]).size().reset_index(name="n")
    g = g.sort_values(["n", "_api_label"], ascending=[False, True], kind="stable")
    return [(a, int(n), c) for a, c, n in zip(g["_api_label"], g["_cat"], g["n"])]


def _by_col(df: pd.DataFrame, col: str) -> list[Row]:
    if df.empty:
        return []
    counts = df[col].value_counts(sort=True)
    ordered = sorted(counts.items(), key=lambda kv: (-kv[1], str(kv[0])))
    return [(str(k), int(v), C.CAT_OTHER) for k, v in ordered]


def _write(ws, col: int, row: int, title: str, headers: tuple[str, str], rows: list[Row], mode: str) -> int:
    """Write a table; return the next free row below it."""
    ws.cell(row, col, title).font = Font(bold=True, size=11)
    fill = {"yellow": YELLOW, "pink": PINK}.get(mode, HEADER_FILL)
    font = Font(bold=True, color="000000" if mode in ("yellow", "pink") else "FFFFFF")
    for i, h in enumerate(headers):
        c = ws.cell(row + 1, col + i, h)
        c.fill, c.font, c.border = fill, font, BOX
    r = row + 2
    for label, n, cat in rows:
        a, b = ws.cell(r, col, label), ws.cell(r, col + 1, n)
        for c in (a, b):
            c.border = BOX
            if mode == "api" and cat in CATEGORY_STYLE:
                f, color = CATEGORY_STYLE[cat]
                c.fill, c.font = f, Font(color=color)
        b.number_format = "#,##0"
        r += 1
    total = ws.cell(r, col + 1, sum(n for _, n, _ in rows))
    total.number_format = "#,##0"
    ws.cell(r, col, "Grand Total").font = Font(bold=True)
    total.font = Font(bold=True)
    for c in (ws.cell(r, col), total):
        c.border = BOX
    return r + 3


def build_pivot(ws, export: pd.DataFrame) -> None:
    cat, bt = export["_cat"], export[C.BOOK_TYPE_COL]
    change = export[cat.isin(C.VIEW_CATEGORIES[C.VIEW_CHANGE])]
    evc = export[cat.isin({C.CAT_EVC, C.CAT_EVC_OPEN})]
    dev = export[cat == C.CAT_DEV]
    dom = C.DOMAIN_COL
    hdr = ("API Result", "Count")

    _write(ws, 1, 1, "Table 1 - All Counts", hdr, _by_api(export), "api")
    nxt = _write(ws, 4, 1, "Table 2 - All PDF Backlog", hdr, _by_api(export[bt == C.BT_PDF]), "api")
    nxt = _write(ws, 4, nxt, "Table 3 - HTML+PDF Backlog", hdr, _by_api(export[bt == C.BT_HTML_PDF]), "api")
    _write(ws, 4, nxt, "Table 4 - HTML Backlog", hdr, _by_api(export[bt == C.BT_HTML]), "api")

    nxt = _write(ws, 7, 1, "Table 5 - Change/BAU by Book Type", (C.BOOK_TYPE_COL, "Count"),
                 _by_col(change, C.BOOK_TYPE_COL), "yellow")
    _write(ws, 7, nxt, "Table 6 - Change/BAU by Domain", (dom, "Count"), _by_col(change, dom), "yellow")
    _write(ws, 10, 1, "Table 7 - EVC by Domain", (dom, "Count"), _by_col(evc, dom), "pink")
    _write(ws, 13, 1, "Table 8 - Stream 2 by Domain", (dom, "Count"), _by_col(dev, dom), "plain")

    widths = {1: 46, 2: 10, 4: 46, 5: 10, 7: 40, 8: 10, 10: 36, 11: 10, 13: 36, 14: 10}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.sheet_view.showGridLines = False
