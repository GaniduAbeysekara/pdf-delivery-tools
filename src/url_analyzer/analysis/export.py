"""The one Excel exporter for URL analysis results (paste, Excel upload, Daily Backlog and CLI all use it)."""
from __future__ import annotations

import re
import time
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import URLAnalysisResult

_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
HEADER_FILL = PatternFill("solid", fgColor="4A6FA5")

# The standard columns, identical for every input source.
STANDARD_COLUMNS = [
    ("document_id", "Document ID"), ("book_title", "Book Title"), ("url", "URL"), ("domain", "Domain"),
    ("api_result", "API Result"), ("book_type", "Book Type"),        # shown only when the source provides them
    ("http_status", "HTTP Status"), ("final_url", "Final URL"), ("content_type", "Content Type"),
    ("direct_document", "Direct Document"), ("page_count", "Page Count"), ("file_size", "File Size"),
    ("language", "Language"), ("last_modified", "Last Modified"), ("analysis_status", "Analysis Status"),
    ("error", "Error"),
]
OPTIONAL_STANDARD = {"api_result", "book_type"}
DETAIL_COLUMNS = [
    ("status", "Status"), ("url_type", "URL Type"), ("document_url", "Document URL"),
    ("document_link_type", "Document Link Type"), ("document_link_count", "Document Link Count"),
    ("redirects", "Redirects"), ("final_status", "Final Status"), ("mime_type", "MIME Type"),
    ("file_size_bytes", "File Size (bytes)"), ("text_extraction_status", "Text Extraction Status"),
    ("cached", "Cached"), ("analyzed_at", "Analyzed At"), ("notes", "Notes"),
]


def _safe(v):
    if isinstance(v, str):
        v = _ILLEGAL.sub("", v)
        if v.startswith("="):          # never let remote or pasted text become an Excel formula
            v = "'" + v
    return v


def summarize(results: list[URLAnalysisResult]) -> list[tuple[str, int]]:
    """The summary numbers shown on the dashboard cards and on the SUMMARY sheet (single definition)."""
    def n(pred) -> int:
        return sum(1 for r in results if pred(r))
    return [
        ("Total URLs", len(results)),
        ("Working", n(lambda r: r.status == "WORKING")),
        ("Not Working", n(lambda r: r.status != "WORKING")),
        ("   Blocked (401/403/429)", n(lambda r: r.status == "BLOCKED")),
        ("   Timeout", n(lambda r: r.status == "TIMEOUT")),
        ("   Error (DNS / connection / SSL / invalid)", n(lambda r: r.status == "ERROR")),
        ("Direct PDF", n(lambda r: r.url_type == "DIRECT_PDF")),
        ("Direct Document", n(lambda r: r.url_type == "DIRECT_DOCUMENT")),
        ("Landing Page", n(lambda r: r.url_type == "LANDING_PAGE")),
        ("Single Document Link", n(lambda r: r.document_link_type == "SINGLE")),
        ("Multiple Document Links", n(lambda r: r.document_link_type == "MULTIPLE")),
        ("No Document Link", n(lambda r: r.document_link_type == "NONE")),
        ("Unknown", n(lambda r: r.url_type == "UNKNOWN")),
        ("Total Pages", sum(int(r.page_count) for r in results if r.page_count.isdigit())),
    ]


def results_frame(results: list[URLAnalysisResult], source_columns: list[str] | None = None) -> pd.DataFrame:
    rows = [r.to_dict() for r in results]
    cols = [(k, label) for k, label in STANDARD_COLUMNS
            if k not in OPTIONAL_STANDARD or any(r.get(k) for r in rows)]
    cols += DETAIL_COLUMNS
    if source_columns is None:
        source_columns = sorted({k for r in rows for k in r["extra"]})
    taken = {label for _, label in cols}
    src = [(c, c if c not in taken else f"{c} (source)") for c in source_columns]
    data = []
    for r in rows:
        row = {}
        for k, label in cols:
            v = r.get(k)
            if k == "notes":
                v = "; ".join(v)
            elif k == "cached":
                v = "Yes" if v else "No"
            elif k == "page_count" and str(v).isdigit():
                v = int(v)
            row[label] = v
        for orig, label in src:
            row[label] = r["extra"].get(orig, "")
        data.append(row)
    df = pd.DataFrame(data, columns=[l for _, l in cols] + [l for _, l in src])
    for c in df.columns:
        df[c] = df[c].map(_safe)
    # real dates where known (the text "Unknown" stays as text)
    df["Last Modified"] = df["Last Modified"].map(
        lambda v: pd.to_datetime(v).to_pydatetime() if isinstance(v, str) and re.match(r"^\d{4}-\d{2}-\d{2}$", v) else v)
    df["Analyzed At"] = df["Analyzed At"].map(
        lambda v: pd.to_datetime(v).to_pydatetime() if isinstance(v, str) and v else None)
    return df


def export_analysis_results(results: list[URLAnalysisResult], path: str | Path | None = None, *,
                            source: str = "", source_name: str = "",
                            source_columns: list[str] | None = None) -> Path:
    """Write results to a formatted workbook and return its path. Same function for every input source."""
    path = Path(path) if path else Path("output") / f"URL_Analysis_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    df = results_frame(results, source_columns)
    with pd.ExcelWriter(path, engine="openpyxl", datetime_format="yyyy-mm-dd hh:mm", date_format="yyyy-mm-dd") as w:
        df.to_excel(w, sheet_name="URL Analysis", index=False)
        ws = w.sheets["URL Analysis"]
        ws.freeze_panes = "A2"
        for i, col in enumerate(df.columns, 1):
            sample = df[col].head(100).map(lambda v: len(str(v)) if v is not None else 0).max() if len(df) else 0
            ws.column_dimensions[get_column_letter(i)].width = max(len(str(col)) + 2, min(int(sample or 0) + 2, 50))
            ws.cell(1, i).alignment = Alignment(vertical="center")
            if col == "Last Modified":
                for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    row[0].number_format = "yyyy-mm-dd"
        if len(df):
            t = Table(displayName="URLAnalysis", ref=f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}")
            t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(t)
        links = [{"Document ID": r.document_id, "Book Title": r.book_title, "Landing Page URL": r.url,
                  "Document Link Type": r.document_link_type, "Document URL": u}
                 for r in results if r.document_link_type in ("SINGLE", "MULTIPLE") for u in r.document_urls]
        if links:
            ldf = pd.DataFrame(links)
            for c in ldf.columns:
                ldf[c] = ldf[c].map(_safe)
            ldf.to_excel(w, sheet_name="Document Links", index=False)
            for col, width in zip("ABCDE", (16, 40, 60, 20, 80)):
                w.sheets["Document Links"].column_dimensions[col].width = width
            w.sheets["Document Links"].freeze_panes = "A2"
        summary = [(k, v) for k, v in summarize(results)]
        info = [("Source", source or "Unspecified"), ("Source file", source_name or "-"),
                ("Generated (UTC)", time.strftime("%Y-%m-%d %H:%M:%S", time.gmtime())),
                ("Results reused from cache", sum(r.cached for r in results))]
        pd.DataFrame(summary + info, columns=["Metric", "Value"]).to_excel(w, sheet_name="SUMMARY", index=False)
        w.sheets["SUMMARY"].column_dimensions["A"].width = 44
        w.sheets["SUMMARY"].column_dimensions["B"].width = 34
        for c in w.sheets["SUMMARY"][1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
    return path
