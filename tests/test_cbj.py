"""Superfinanciera CBJ crawler - the logic that needs no browser: numerals, sidebar parsing, de-duplication, output files."""
import json

import pytest
from openpyxl import load_workbook

from url_analyzer.recipes import cbj
from url_analyzer.recipes.cbj import (Attachment, ChapterRun, Result, absolute, dedupe, file_ext, ordinal_of, parse_sidebar, roman_to_int,
                                      tracker_script, unique_urls, write_outputs)

P1, T1, C3 = "Parte I. Instrucciones", "Título I. Aspectos generales", "Capítulo III. Gobierno corporativo"
URL = "https://www.superfinanciera.gov.co/sfcservices/api-circular-basica/api-circular-basica/apiCircularBasica/descargarAnexo/328"


def att(**kw):
    base = dict(parte=P1, titulo=T1, capitulo=C3, name="Anexo 1.xls", url=URL, ext="xls", found_in="api")
    return Attachment(**{**base, **kw})


# ── numerals and labels ──────────────────────────────────────────────────────
@pytest.mark.parametrize("roman,n", [("I", 1), ("IV", 4), ("IX", 9), ("X", 10), ("XIV", 14), ("XL", 40), ("xii", 12), ("7", 7)])
def test_roman_numerals(roman, n):
    assert roman_to_int(roman) == n


def test_not_a_numeral():
    with pytest.raises(ValueError):
        roman_to_int("ABC")


@pytest.mark.parametrize("label,n", [("Capítulo III. Gobierno corporativo", 3), ("Parte II. Mercado intermediado", 2), ("Título X. Algo", 10),
                                     ("Capitulo 5. Sin tilde", 5), ("Anexo 1", None), ("", None)])
def test_ordinal_of_a_label(label, n):
    assert ordinal_of(label) == n


# ── the sidebar ──────────────────────────────────────────────────────────────
def row(i, text, rid="", btn=None, rotated=None, selected=False):
    return {"i": i, "id": rid, "text": text, "btn": btn, "rotated": rotated, "selected": selected}


SIDEBAR = [row(0, "Parte I. Uno", btn="flechaParte", rotated=True), row(1, "Título I. A", "titulos", "flecha", True), row(2, "Capítulo I. X", "capitulos", selected=True),
           row(3, "Capítulo II. Y", "capitulos"), row(4, "Título II. B", "titulos", "flecha", False),
           row(5, "Parte II. Dos", btn="flechaParte", rotated=True), row(6, "Título III. Repetido", "titulos", "flecha", True), row(7, "Capítulo I. P", "capitulos"),
           row(8, "Título III. Repetido", "titulos", "flecha", True), row(9, "Capítulo I. Q", "capitulos"), row(10, "  ", "capitulos")]


def test_sidebar_rows_become_the_hierarchy_of_every_capitulo():
    got = parse_sidebar(SIDEBAR)
    assert [(c["parte"], c["titulo"], c["capitulo"]) for c in got] == [
        ("Parte I. Uno", "Título I. A", "Capítulo I. X"), ("Parte I. Uno", "Título I. A", "Capítulo II. Y"),
        ("Parte II. Dos", "Título III. Repetido", "Capítulo I. P"), ("Parte II. Dos", "Título III. Repetido", "Capítulo I. Q")]   # the blank row is ignored
    assert [c["row"] for c in got] == [2, 3, 7, 9]


def test_a_title_label_listed_twice_under_one_parte_is_told_apart():
    got = parse_sidebar(SIDEBAR)
    assert [c["titulo_occ"] for c in got] == [0, 0, 0, 1]                 # the second "Título III. Repetido" is occurrence 1
    assert cbj.CbjCrawler._titulo_row(SIDEBAR, "Parte II. Dos", "Título III. Repetido", 0)["i"] == 6
    assert cbj.CbjCrawler._titulo_row(SIDEBAR, "Parte II. Dos", "Título III. Repetido", 1)["i"] == 8
    assert cbj.CbjCrawler._titulo_row(SIDEBAR, "Parte II. Dos", "Título III. Repetido", 2) is None
    assert cbj.CbjCrawler._titulo_row(SIDEBAR, "Parte I. Uno", "Título III. Repetido") is None            # other Parte


# ── files and addresses ──────────────────────────────────────────────────────
@pytest.mark.parametrize("name,url,ct,ext", [
    ("Anexo 1 Formato Revisor Fiscal.xls", "", "", "xls"), ("Anexo.DOCX", "", "", "docx"), ("Guia.pdf ", "", "", "pdf"),
    ("sin extension", URL, "application/vnd.ms-excel", "xls"), ("", "https://a.org/x/y.PDF?v=1", "", "pdf"), ("", "https://a.org/x/y.pdf", "", "pdf"),
    ("x", "", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"), ("x", "", "application/octet-stream", ""),
    ("reporte.final.xlsx", "", "", "xlsx")])
def test_file_type(name, url, ct, ext):
    assert file_ext(name, url, ct) == ext


def test_relative_addresses_become_absolute():
    assert absolute("https://s.gov/visualizador/inicio", "../files/a.pdf") == "https://s.gov/files/a.pdf"
    assert absolute("https://s.gov/a/b", "/c.pdf") == "https://s.gov/c.pdf"
    assert absolute("https://s.gov/a/b", "//cdn.org/c.pdf") == "https://cdn.org/c.pdf"
    assert absolute("https://s.gov/a/b", "") == ""


# ── duplicates without losing the hierarchy ──────────────────────────────────
def test_the_same_file_twice_in_one_capitulo_is_one_row_with_what_each_sighting_knew():
    rows = dedupe([att(found_in="api", content_type="application/vnd.ms-excel", attachment_id="328"),
                   att(found_in="click", verified_by_click=True), att(found_in="dom", name="")])
    assert len(rows) == 1
    r = rows[0]
    assert (r.found_in, r.verified_by_click, r.content_type, r.attachment_id, r.name) == ("api, click, dom", True, "application/vnd.ms-excel", "328", "Anexo 1.xls")


def test_the_same_file_in_two_capitulos_keeps_both_rows_and_both_places():
    rows = dedupe([att(), att(capitulo="Capítulo IV. Otro"), att(title_row=2)])
    assert [(r.capitulo, r.title_row) for r in rows] == [(C3, 1), ("Capítulo IV. Otro", 1), (C3, 2)]
    u = unique_urls(rows)
    assert len(u) == 1 and u[0]["occurrences"] == 3 and u[0]["locations"].count(" | ") == 2
    assert "[title row 2]" in u[0]["locations"]


def test_what_was_learnt_by_reading_the_file_is_kept_when_rows_are_merged():
    rows = dedupe([att(found_in="api"), att(found_in="click", size_bytes=2048, pages="7", language="English", last_modified="2025-10-01")])
    assert len(rows) == 1 and (rows[0].size_bytes, rows[0].pages, rows[0].language, rows[0].last_modified) == (2048, "7", "English", "2025-10-01")


def test_a_numeral_level_sighting_upgrades_the_level():
    rows = dedupe([att(level="Capitulo"), att(level="Numeral", numeral="2")])
    assert rows[0].level == "Numeral" and rows[0].numeral == "2"


def test_attachments_without_an_address_are_still_kept_by_name():
    rows = dedupe([att(url="", name="Anexo A.xls"), att(url="", name="Anexo A.xls"), att(url="", name="Anexo B.xls")])
    assert [r.name for r in rows] == ["Anexo A.xls", "Anexo B.xls"]


# ── the wait-condition tracker ───────────────────────────────────────────────
def test_tracker_script_counts_only_the_sites_own_requests():
    js = tracker_script(r"/sfcservices/")
    assert 'new RegExp("/sfcservices/")' in js and "XMLHttpRequest.prototype.send" in js and "window.fetch" in js
    assert "googleapis" not in js and "firestore" not in js.lower()                    # third-party long-lived connections are not counted
    assert "inflight" in cbj.JS_IDLE


# ── output ───────────────────────────────────────────────────────────────────
def test_output_workbook_and_json(tmp_path):
    res = Result([att(attachment_id="328", content_type="application/vnd.ms-excel", found_in="api, click", verified_by_click=True, app_ids="1/1/3"),
                  att(capitulo="Capítulo IV. Otro", app_ids="1/1/4"), att(name="=cmd|' /C calc'!A0", url="https://x.org/e")],
                 [ChapterRun(P1, T1, C3, status="ok", attachments=1), ChapterRun(P1, T1, "Capítulo V. Roto", status="failed", error="TimeoutError: x")],
                 "2026-10-08 10:00:00", 12.5)
    out, js = write_outputs(res, tmp_path / "sub" / "cbj.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Attachments", "Unique URLs", "Capitulos", "Errors"]
    head = [c.value for c in wb["Attachments"][1]]
    assert head[:3] == ["Parte", "Titulo", "Capitulo"] and "Attachment URL" in head and "Site IDs (Parte/Titulo/Capitulo)" in head
    first = {h: c.value for h, c in zip(head, wb["Attachments"][2])}
    assert first["Attachment URL"] == URL and first["Type"] == "xls" and first["Verified By Click"] is True
    assert wb["Attachments"].max_row == 4 and wb["Unique URLs"].max_row == 3                       # 3 rows; the same URL in two Capitulos is one unique URL
    assert [c.value for c in wb["Errors"][2]][3:5] == ["failed", 0] or wb["Errors"].max_row == 2
    assert any(str(c.value).startswith("'=") for row in wb["Attachments"].iter_rows(min_row=2) for c in row)          # formula injection guard
    data = json.loads(js.read_text(encoding="utf-8"))
    assert data["summary"]["Capitulos failed"] == 1 and len(data["attachments"]) == 3
    assert data["attachments"][0]["hierarchy"] == f"{P1} > {T1} > {C3}"
    assert data["chapters"][1]["error"] == "TimeoutError: x"


def test_result_lists_failed_chapters():
    r = Result(chapters=[ChapterRun("p", "t", "c1", status="ok"), ChapterRun("p", "t", "c2", status="failed")])
    assert [c.capitulo for c in r.errors] == ["c2"]
