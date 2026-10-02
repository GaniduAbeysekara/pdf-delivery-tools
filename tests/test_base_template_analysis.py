"""Base Template Analysis: dataset vs Base Templates joined on domain, flags, report, HTTP endpoints."""
import io
import sys
from pathlib import Path

import pandas as pd
import pytest
from flask import Flask
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

from url_analyzer.base_templates import compare as C
from url_analyzer.base_templates.analysis_web import create_analysis_blueprint
from url_analyzer.base_templates.compare_report import build_report
from url_analyzer.base_templates.store import BaseTemplateStore
from url_analyzer.base_templates.web import create_base_templates_blueprint

H = {"X-Requested-With": "fetch"}
MAP = {"document_id": "DocID", "domain": "Domain", "spidering_template": "SpiderTemplate", "url": "URL"}
BASE = [
    {"id": "1", "domain": "www.a.org", "template": "BASE-A", "created_date": "2026-09-01", "dev_name": "Dev"},
    {"id": "2", "domain": "b.org, c.org", "template": "BASE-BC", "created_date": "", "dev_name": ""},
    {"id": "3", "domain": "", "template": "BASE-NODOM", "created_date": "", "dev_name": ""},
    {"id": "4", "domain": "unused.org", "template": "BASE-UNUSED", "created_date": "", "dev_name": ""},
]


def frame(rows):
    return pd.DataFrame(rows, columns=["DocID", "Domain", "SpiderTemplate", "URL"])


def run(rows, base=BASE):
    return C.compare(frame(rows), MAP, base)


def by_doc(res):
    return {r["document_id"]: r for r in res.rows}


# ── flags ────────────────────────────────────────────────────────────────────
def test_clean_row_is_ok_and_joined():
    r = by_doc(run([["D1", "www.a.org", "TPL", "https://www.a.org/x.pdf"]]))["D1"]
    assert r["flags"] == [] and r["status"] == "OK" and r["base_templates"] == ["BASE-A"] and r["match_type"] == "Exact"


def test_missing_url_and_template_flagged():
    r = by_doc(run([["D1", "www.a.org", "", ""], ["D2", "www.a.org", "T", "  "]]))
    assert r["D1"]["flags"] == ["MISSING_URL", "MISSING_SPIDER_TEMPLATE"] and r["D1"]["status"] == "Needs fix"
    assert r["D2"]["flags"] == ["MISSING_URL"]


def test_placeholders_count_as_missing():
    r = by_doc(run([["D1", "www.a.org", "N/A", "nan"]]))["D1"]
    assert set(r["flags"]) == {"MISSING_URL", "MISSING_SPIDER_TEMPLATE"}


def test_domain_not_in_base_templates():
    r = by_doc(run([["D1", "nowhere.org", "T", "https://nowhere.org/a"]]))["D1"]
    assert r["flags"] == ["NO_BASE_TEMPLATE"] and r["base_templates"] == []


def test_domain_is_normalised_before_joining():
    r = by_doc(run([["D1", "HTTPS://WWW.A.ORG/path/", "T", "u"]]))["D1"]
    assert r["domain"] == "www.a.org" and r["flags"] == []


def test_cached_domain_with_several_domains_joins_each():
    res = by_doc(run([["D1", "b.org", "T", "u"], ["D2", "c.org", "T", "u"]]))
    assert res["D1"]["base_templates"] == res["D2"]["base_templates"] == ["BASE-BC"]


def test_www_only_difference_matches_with_info_flag():
    r = by_doc(run([["D1", "a.org", "T", "u"], ["D2", "www.b.org", "T", "u"]]))
    assert r["D1"]["flags"] == ["WWW_ONLY_MATCH"] and r["D1"]["status"] == "Check" and r["D1"]["base_templates"] == ["BASE-A"]
    assert r["D2"]["match_type"] == "Ignoring www"


def test_domain_taken_from_url_when_blank_and_missing_domain_flag():
    res = by_doc(run([["D1", "", "T", "https://www.a.org/x"], ["D2", "", "T", ""], ["D3", "Unknown", "T", "not a url"]]))
    assert res["D1"]["domain"] == "www.a.org" and res["D1"]["domain_from_url"] and res["D1"]["flags"] == []
    assert "MISSING_DOMAIN" in res["D2"]["flags"] and "MISSING_URL" in res["D2"]["flags"] and "NO_BASE_TEMPLATE" not in res["D2"]["flags"]
    assert "MISSING_DOMAIN" in res["D3"]["flags"]


def test_base_side_flags_and_document_counts():
    res = run([["D1", "www.a.org", "T", "u"], ["D2", "www.a.org", "T", "u"], ["D3", "b.org", "T", "u"]])
    base = {b["template"]: b for b in res.base}
    assert base["BASE-A"]["documents"] == 2 and base["BASE-A"]["flags"] == []
    assert base["BASE-NODOM"]["flags"] == ["BASE_NO_CACHED_DOMAIN"]
    assert base["BASE-UNUSED"]["flags"] == ["BASE_NOT_USED"]
    assert base["BASE-BC"]["documents"] == 1


def test_non_address_cached_domain_text_cannot_join():
    res = run([["D1", "www.a.org", "T", "u"]], [{"id": "9", "domain": "Home | BCSC", "template": "X"}])
    assert res.base[0]["flags"] == ["BASE_NO_CACHED_DOMAIN"]


def test_domain_summary_and_totals():
    res = run([["D1", "www.a.org", "T", "u"], ["D2", "www.a.org", "", "u"], ["D3", "zzz.org", "T", ""]])
    dom = {d["domain"]: d for d in res.domains}
    assert dom["www.a.org"]["rows"] == 2 and dom["www.a.org"]["needs_fix"] == 1 and dom["www.a.org"]["missing_template"] == 1
    assert dom["zzz.org"]["status"] == "No Base Template"
    s = res.summary
    assert (s["dataset_rows"], s["rows_ok"], s["rows_needs_fix"]) == (3, 1, 2)
    assert s["flags"]["MISSING_URL"] == 1 and s["flags"]["NO_BASE_TEMPLATE"] == 1 and s["domains_without_base"] == 1


def test_filter_rows():
    res = run([["D1", "www.a.org", "T", "u"], ["D2", "zzz.org", "T", "u"]])
    assert [r["document_id"] for r in C.filter_rows(res.rows, flag="NO_BASE_TEMPLATE")] == ["D2"]
    assert [r["document_id"] for r in C.filter_rows(res.rows, status="OK")] == ["D1"]
    assert [r["document_id"] for r in C.filter_rows(res.rows, q="base-a t")] == ["D1"]
    assert [r["document_id"] for r in C.filter_rows(res.rows, domain="zzz.org")] == ["D2"]


def test_mapping_validation():
    with pytest.raises(C.CompareError, match="URL"):
        C.check_mapping(["a", "b"], {"spidering_template": "a"})
    with pytest.raises(C.CompareError, match="Spider Template"):
        C.check_mapping(["a", "b"], {"url": "a"})
    with pytest.raises(C.CompareError, match="same column"):
        C.check_mapping(["a"], {"url": "a", "spidering_template": "a"})
    with pytest.raises(C.CompareError, match="does not exist"):
        C.check_mapping(["a", "b"], {"url": "x", "spidering_template": "b"})


# ── reading files ────────────────────────────────────────────────────────────
def write_xlsx(path, rows, header=("Document ID", "Domain", "SpideringTemplate", "URL"), title_rows=0, sheet="URL Analysis"):
    wb = Workbook()
    wb.active.title = "Other"
    ws = wb.create_sheet(sheet)
    for _ in range(title_rows):
        ws.append(["Exported report"])
    ws.append(list(header))
    for r in rows:
        ws.append(r)
    wb.move_sheet("Other", offset=1)
    wb.save(path)
    return path


def test_load_xlsx_detects_columns_row_numbers_and_skips_blank_rows(tmp_path):
    p = write_xlsx(tmp_path / "a.xlsx", [["D1", "www.a.org", "T", "u"], [None] * 4, ["D2", "", "", ""]], title_rows=2)
    ds = C.load_dataset(p)
    assert ds.sheet == "URL Analysis" and ds.header_row == 3
    assert ds.mapping == {"url": "URL", "spidering_template": "SpideringTemplate", "domain": "Domain", "document_id": "Document ID"}
    assert ds.row_numbers == [4, 6] and len(ds.frame) == 2           # the empty row 5 is ignored; numbers match the sheet
    res = C.compare(ds, ds.mapping, BASE)
    assert [r["source_row"] for r in res.rows] == [4, 6]


def test_load_csv_with_semicolons_and_spider_template_heading(tmp_path):
    p = tmp_path / "a.csv"
    p.write_text("DocID;Domain;Spider Template;URL\nD1;www.a.org;T;http://www.a.org/x\n", encoding="utf-8")
    ds = C.load_dataset(p)
    assert ds.mapping["spidering_template"] == "Spider Template" and ds.frame.iloc[0]["URL"] == "http://www.a.org/x"


def test_unreadable_file_and_unknown_sheet(tmp_path):
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"nope")
    with pytest.raises(C.CompareError):
        C.load_dataset(bad)
    with pytest.raises(C.CompareError, match="not found"):
        C.load_dataset(write_xlsx(tmp_path / "ok.xlsx", [["D", "d.org", "T", "u"]]), sheet="Missing")


# ── report ───────────────────────────────────────────────────────────────────
def test_report_sheets_flags_and_subset():
    res = run([["D1", "www.a.org", "T", "u"], ["D2", "zzz.org", "", ""]])
    wb = load_workbook(io.BytesIO(build_report(res)))
    assert wb.sheetnames == ["SUMMARY", "Review", "Needs Fix", "Domains", "Base Templates Review"]
    ws = wb["Review"]
    head = [c.value for c in ws[1]]
    assert head[:7] == ["Source Row", "Document ID", "Domain", "Spider Template", "URL", "Status", "Flags"]
    row2 = {h: c.value for h, c in zip(head, ws[3])}
    assert row2["Status"] == "Needs fix" and row2["Missing URL"] == "YES" and row2["No Base Template"] == "YES"
    assert [c.value for c in wb["Needs Fix"]["B"]][1:] == ["D2"]
    only = load_workbook(io.BytesIO(build_report(res, [res.rows[0]])))
    assert [c.value for c in only["Review"]["B"]][1:] == ["D1"]
    assert any("Needs fix" == c.value for c in wb["SUMMARY"]["A"])


def test_report_guards_formula_injection():
    res = run([["D1", "www.a.org", "=1+1", "=HYPERLINK(\"x\")"]])
    ws = load_workbook(io.BytesIO(build_report(res)))["Review"]
    assert str(ws["D2"].value).startswith("'=")


# ── HTTP ─────────────────────────────────────────────────────────────────────
@pytest.fixture
def client(tmp_path):
    app = Flask(__name__, template_folder=str(ROOT / "templates"))
    store = BaseTemplateStore(tmp_path / "bt.json")
    app.register_blueprint(create_base_templates_blueprint(tmp_path / "bt.json", store))
    app.register_blueprint(create_analysis_blueprint(tmp_path / "up", store))
    app.testing = True
    store.add({"template": "BASE-A", "domain": "https://www.a.org/", "created_date": "2026-09-01"})
    store.add({"template": "BASE-B", "domain": "b.org"})
    return app.test_client()


def upload(client, tmp_path, rows):
    p = write_xlsx(tmp_path / "ds.xlsx", rows)
    with open(p, "rb") as fh:
        return client.post("/api/bta/upload", data={"file": (fh, "ds.xlsx")}, headers=H)


def test_page_renders(client):
    html = client.get("/base-template-analysis").data.decode()
    assert "Base Template Analysis" in html and "No Base Template" in html


def test_upload_run_filter_download(client, tmp_path):
    rows = [["D1", "www.a.org", "T", "https://www.a.org/1"], ["D2", "b.org", "", "https://b.org/2"],
            ["D3", "nope.org", "T", ""], ["D4", "www.a.org", "T", "https://www.a.org/4"]]
    up = upload(client, tmp_path, rows)
    assert up.status_code == 200
    d = up.get_json()
    assert d["rows"] == 4 and d["mapping"]["url"] == "URL" and d["sheet"] == "URL Analysis"
    r = client.post("/api/bta/run", json={"upload_id": d["upload_id"], "sheet": d["sheet"], "mapping": d["mapping"]}, headers=H)
    assert r.status_code == 200
    out = r.get_json()
    rid, s = out["result_id"], out["summary"]
    assert (s["dataset_rows"], s["rows_ok"], s["rows_needs_fix"], s["base_rows"]) == (4, 2, 2, 2)
    lst = client.get(f"/api/bta/{rid}/rows?flag=NO_BASE_TEMPLATE").get_json()
    assert [i["document_id"] for i in lst["items"]] == ["D3"]
    assert client.get(f"/api/bta/{rid}/rows?view=domains").get_json()["matched"] == 3
    assert client.get(f"/api/bta/{rid}/rows?view=base&flag=BASE_NOT_USED").get_json()["matched"] == 0
    assert client.get(f"/api/bta/{rid}/rows?page=x").status_code == 400
    dl = client.get(f"/api/bta/{rid}/download?scope=fix")
    assert dl.status_code == 200 and dl.data[:2] == b"PK"
    ws = load_workbook(io.BytesIO(dl.data))["Review"]
    assert {c.value for c in ws["B"]} - {None} == {"Document ID", "D2", "D3"}


def test_uses_latest_base_templates_each_run(client, tmp_path):
    d = upload(client, tmp_path, [["D3", "nope.org", "T", "u"]]).get_json()
    body = {"upload_id": d["upload_id"], "sheet": d["sheet"], "mapping": d["mapping"]}
    first = client.post("/api/bta/run", json=body, headers=H).get_json()["summary"]
    assert first["flags"]["NO_BASE_TEMPLATE"] == 1
    client.post("/api/base-templates", json={"template": "BASE-NEW", "domain": "nope.org"}, headers=H)
    second = client.post("/api/bta/run", json=body, headers=H).get_json()["summary"]
    assert second["flags"]["NO_BASE_TEMPLATE"] == 0 and second["rows_ok"] == 1


def test_errors_are_user_facing(client, tmp_path):
    assert client.post("/api/bta/upload", data={}, headers=H).status_code == 400
    assert client.post("/api/bta/upload", data={"file": (io.BytesIO(b"x"), "x.txt")}, headers=H).status_code == 400
    assert client.post("/api/bta/upload", data={"file": (io.BytesIO(b"x"), "x.xlsx")}, headers=H).status_code == 400
    assert client.post("/api/bta/upload", data={}).status_code == 400                      # missing header
    assert client.post("/api/bta/run", json={"upload_id": "nope"}, headers=H).status_code == 404
    assert client.get("/api/bta/zzz/rows").status_code == 404
    d = upload(client, tmp_path, [["D", "d.org", "T", "u"]]).get_json()
    r = client.post("/api/bta/run", json={"upload_id": d["upload_id"], "mapping": {"url": "URL"}}, headers=H)
    assert r.status_code == 400 and "Spider Template" in r.get_json()["error"]


def test_nav_and_home_in_full_app():
    import importlib
    dash = importlib.import_module("dashboard")
    c = dash.create_app().test_client()
    assert 'href="/base-template-analysis"' in c.get("/").data.decode()
    html = c.get("/base-template-analysis").data.decode()
    assert 'class="nav on" href="/base-template-analysis"' in html and 'class="nav on" href="/base-templates"' not in html


# ── hand-off from the URL Analysis page ──────────────────────────────────────
from tests.test_shared_engine import U1, U2, U3, env, wait_job  # noqa: E402,F401  (fixtures reused: fake web + analysis service)
from url_analyzer.ui.url_analysis import create_url_analysis_blueprint  # noqa: E402


@pytest.fixture
def both(env, tmp_path):
    store = BaseTemplateStore(tmp_path / "bt.json")
    store.add({"template": "BASE-EX", "domain": "https://example.gov/"})
    app = Flask(__name__, template_folder=str(ROOT / "templates"))
    app.register_blueprint(create_url_analysis_blueprint(tmp_path / "out2", tmp_path / "up2", service=env.svc))
    app.register_blueprint(create_base_templates_blueprint(tmp_path / "bt.json", store))
    app.register_blueprint(create_analysis_blueprint(tmp_path / "bta_up", store))
    app.testing = True
    return app.test_client(), store


def finished_job(client):
    items = [{"url": U1, "document_id": "A1", "domain": "example.gov", "spidering_template": "TPL-1"},
             {"url": U2, "document_id": "A2", "domain": "other.org", "spidering_template": ""},
             {"url": U3, "document_id": "A3", "domain": "example.gov", "spidering_template": "TPL-3"}]
    jid = client.post("/api/url-analysis/analyze", json={"mode": "items", "items": items}).get_json()["job_id"]
    wait_job(client, jid)
    return jid


def test_send_all_and_selected_from_url_analysis(both):
    client, _ = both
    jid = finished_job(client)
    all_ = client.post(f"/api/url-analysis/{jid}/send-to-base-templates", json={}).get_json()
    assert all_["count"] == 3
    sel = client.post(f"/api/url-analysis/{jid}/send-to-base-templates", json={"indices": [1]})
    assert sel.get_json()["count"] == 1
    from url_analyzer import transfer
    rec = transfer.get(sel.get_json()["transfer_id"])["records"][0]
    assert rec == {"row": 2, "document_id": "A2", "domain": "other.org", "spidering_template": "", "url": U2}


def test_send_validation(both):
    client, _ = both
    jid = finished_job(client)
    assert client.post(f"/api/url-analysis/{jid}/send-to-base-templates", json={"indices": []}).status_code == 400
    assert client.post(f"/api/url-analysis/{jid}/send-to-base-templates", json={"indices": ["x"]}).status_code == 400
    assert client.post(f"/api/url-analysis/{jid}/send-to-base-templates", json={"indices": [99]}).status_code == 409
    assert client.post("/api/url-analysis/nope/send-to-base-templates", json={}).status_code == 404


def test_transfer_is_compared_with_base_templates(both):
    client, store = both
    jid = finished_job(client)
    tid = client.post(f"/api/url-analysis/{jid}/send-to-base-templates", json={}).get_json()["transfer_id"]
    out = client.post("/api/bta/from-transfer", json={"transfer_id": tid}, headers=H).get_json()
    s = out["summary"]
    assert out["received"] == 3 and (s["rows_ok"], s["rows_needs_fix"]) == (2, 1)
    rows = {r["document_id"]: r for r in client.get(f"/api/bta/{out['result_id']}/rows").get_json()["items"]}
    assert rows["A1"]["status"] == "OK" and rows["A1"]["base_templates"] == ["BASE-EX"]
    assert rows["A2"]["flags"] == ["MISSING_SPIDER_TEMPLATE", "NO_BASE_TEMPLATE"] and rows["A2"]["source_row"] == 2
    # "Compare again" sees a Base Template added afterwards
    store.add({"template": "BASE-OTHER", "domain": "other.org"})
    again = client.post("/api/bta/from-transfer", json={"transfer_id": tid}, headers=H).get_json()["summary"]
    assert again["flags"]["NO_BASE_TEMPLATE"] == 0
    dl = client.get(f"/api/bta/{out['result_id']}/download")
    assert dl.status_code == 200 and dl.data[:2] == b"PK"


def test_expired_or_missing_transfer(both):
    client, _ = both
    assert client.post("/api/bta/from-transfer", json={"transfer_id": "gone"}, headers=H).status_code == 404
    assert client.post("/api/bta/from-transfer", json={"transfer_id": "gone"}).status_code == 400


def test_button_and_receiver_are_in_the_pages():
    import importlib
    c = importlib.import_module("dashboard").create_app().test_client()
    assert 'id="to-bta"' in c.get("/url-analysis").data.decode()
    assert 'id="recompare"' in c.get("/base-template-analysis?transfer=abc").data.decode()
