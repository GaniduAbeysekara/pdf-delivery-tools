"""Base Template Analysis: by-domain overview (OK / Check / Needs fix per domain), selection, and downloads of selected data."""
import io

import pytest
from flask import Flask
from openpyxl import load_workbook

from tests.test_base_template_analysis import H, ROOT, run, upload
from url_analyzer.base_templates import compare as C
from url_analyzer.base_templates.analysis_web import create_analysis_blueprint
from url_analyzer.base_templates.compare_report import build_report
from url_analyzer.base_templates.store import BaseTemplateStore
from url_analyzer.base_templates.web import create_base_templates_blueprint

ROWS_MIX = [["D1", "www.a.org", "T", "u1"], ["D2", "www.a.org", "", "u2"], ["D3", "www.a.org", "T", ""], ["D4", "a.org", "T", "u4"],
            ["D5", "zzz.org", "T", "u5"], ["D6", "zzz.org", "", ""], ["D7", "b.org", "T", "u7"]]


def ids(ws):
    return [c.value for c in ws["B"]][1:]


# ── the numbers ──────────────────────────────────────────────────────────────
def test_domain_summary_counts_ok_check_and_fix_per_domain():
    res = run(ROWS_MIX)
    dom = {d["domain"]: d for d in res.domains}
    a = dom["www.a.org"]
    assert (a["rows"], a["ok"], a["check"], a["needs_fix"]) == (3, 1, 0, 2)
    assert (a["missing_url"], a["missing_template"], a["status"]) == (1, 1, "Needs fix")
    assert dom["a.org"]["status"] == "Check" and dom["a.org"]["check"] == 1          # only matched after ignoring www
    assert dom["b.org"]["status"] == "OK" and dom["b.org"]["ok"] == 1
    assert (dom["zzz.org"]["status"], dom["zzz.org"]["needs_fix"]) == ("No Base Template", 2)
    for d in res.domains:
        assert d["rows"] == d["ok"] + d["check"] + d["needs_fix"]
    s = res.summary
    assert (s["domains"], s["domains_ok"], s["domains_check"], s["domains_needs_fix"], s["domains_without_base"]) == (4, 1, 1, 1, 1)
    assert s["domains"] == s["domains_ok"] + s["domains_check"] + s["domains_needs_fix"] + s["domains_without_base"] + s["domains_no_domain"]
    assert res.domains[0]["needs_fix"] == max(d["needs_fix"] for d in res.domains)    # worst first


def test_every_row_has_a_stable_index():
    assert [r["idx"] for r in run(ROWS_MIX).rows] == list(range(7))


def test_rows_without_any_domain_form_their_own_group():
    res = run([["D1", "", "T", ""], ["D2", "", "T", ""]])
    assert [(d["domain"], d["rows"], d["status"]) for d in res.domains] == [("(none)", 2, "No domain")]
    assert res.summary["domains_no_domain"] == 1


def test_subset_comparison_recomputes_everything_for_those_rows():
    res = run(ROWS_MIX)
    sub = C.subset_comparison(res, [res.rows[0], res.rows[6]])                        # D1 and D7, both OK
    assert (sub.summary["dataset_rows"], sub.summary["rows_ok"], sub.summary["rows_needs_fix"]) == (2, 2, 0)
    assert {d["domain"] for d in sub.domains} == {"www.a.org", "b.org"} and sub.summary["domains_needs_fix"] == 0
    assert sorted(b["template"] for b in sub.base) == ["BASE-A", "BASE-BC"]           # only the Base Templates these books use
    assert all("BASE_NOT_USED" not in b["flags"] for b in sub.base)


def test_filter_domains():
    res = run(ROWS_MIX)
    assert [d["domain"] for d in C.filter_domains(res.domains, status="No Base Template")] == ["zzz.org"]
    assert [d["domain"] for d in C.filter_domains(res.domains, q="base-bc")] == ["b.org"]


# ── the Excel sheet ──────────────────────────────────────────────────────────
def test_full_report_domains_sheet_has_ok_check_fix_columns():
    wb = load_workbook(io.BytesIO(build_report(run(ROWS_MIX))))
    head = [c.value for c in wb["Domains"][1]]
    assert head[:5] == ["Domain", "Books", "OK", "Check", "Needs fix"]
    lines = {r[0].value: r[1].value for r in wb["SUMMARY"].iter_rows()}
    assert lines["Scope"] == "All books" and lines["Books (documents) in this report"] == 7
    assert lines["Domains OK"] == 1 and lines["Domains with books to fix"] == 1


def test_report_for_a_selection_describes_only_the_selection():
    res = run(ROWS_MIX)
    wb = load_workbook(io.BytesIO(build_report(res, [res.rows[4], res.rows[5]], "Selected domains")))
    assert ids(wb["Review"]) == ["D5", "D6"]
    assert [c.value for c in wb["Domains"]["A"]][1:] == ["zzz.org"]
    head = [c.value for c in wb["Domains"][1]]
    row = {h: c.value for h, c in zip(head, wb["Domains"][2])}
    assert (row["Books"], row["OK"], row["Check"], row["Needs fix"], row["Status"]) == (2, 0, 0, 2, "No Base Template")
    lines = {r[0].value: r[1].value for r in wb["SUMMARY"].iter_rows()}
    assert "2 of 7 books" in lines["Scope"] and lines["Books (documents) in this report"] == 2


# ── HTTP ─────────────────────────────────────────────────────────────────────
@pytest.fixture
def client(tmp_path):
    app = Flask(__name__, template_folder=str(ROOT / "templates"))
    store = BaseTemplateStore(tmp_path / "bt.json")
    for t, d in (("BASE-A", "www.a.org"), ("BASE-B", "b.org")):
        store.add({"template": t, "domain": d})
    app.register_blueprint(create_base_templates_blueprint(tmp_path / "bt.json", store))
    app.register_blueprint(create_analysis_blueprint(tmp_path / "up", store))
    app.testing = True
    return app.test_client()


@pytest.fixture
def report(client, tmp_path):
    d = upload(client, tmp_path, ROWS_MIX).get_json()
    out = client.post("/api/bta/run", json={"upload_id": d["upload_id"], "sheet": d["sheet"], "mapping": d["mapping"]}, headers=H).get_json()
    return out["result_id"], out["summary"]


def test_domains_view_lists_counts_and_filters_by_status(client, report):
    rid, summary = report
    assert summary["domains"] == 4
    items = client.get(f"/api/bta/{rid}/rows?view=domains").get_json()["items"]
    assert {i["domain"]: (i["ok"], i["check"], i["needs_fix"]) for i in items}["www.a.org"] == (1, 0, 2)
    only = client.get(f"/api/bta/{rid}/rows?view=domains&status=No Base Template").get_json()
    assert [i["domain"] for i in only["items"]] == ["zzz.org"]


def test_ids_endpoint_for_select_all_filtered(client, report):
    rid, _ = report
    rows = client.get(f"/api/bta/{rid}/ids?view=rows&status=Needs fix").get_json()
    assert rows["ids"] == [1, 2, 4, 5] and rows["books"] == 4
    doms = client.get(f"/api/bta/{rid}/ids?view=domains&status=Needs fix").get_json()
    assert doms["ids"] == ["www.a.org"] and doms["books"] == 3
    assert client.get(f"/api/bta/{rid}/ids?view=base").status_code == 400
    assert client.get("/api/bta/nope/ids").status_code == 404


def test_download_selected_books_and_domains(client, report):
    rid, _ = report
    url = f"/api/bta/{rid}/download-selected"
    r = client.post(url, json={"rows": [0, 6]}, headers=H)
    assert r.status_code == 200 and r.data[:2] == b"PK"
    assert ids(load_workbook(io.BytesIO(r.data))["Review"]) == ["D1", "D7"]
    r = client.post(url, json={"domains": ["zzz.org", "b.org"]}, headers=H)
    assert ids(load_workbook(io.BytesIO(r.data))["Review"]) == ["D5", "D6", "D7"]          # every book of those domains
    both = client.post(url, json={"rows": [0], "domains": ["b.org"]}, headers=H)
    assert ids(load_workbook(io.BytesIO(both.data))["Review"]) == ["D1", "D7"]


def test_download_selected_validation(client, report):
    rid, _ = report
    url = f"/api/bta/{rid}/download-selected"
    assert client.post(url, json={"rows": [0]}).status_code == 400                         # missing header
    assert client.post(url, json={}, headers=H).status_code == 400                         # nothing selected
    assert client.post(url, json={"rows": ["x"]}, headers=H).status_code == 400
    assert client.post(url, json={"rows": [999]}, headers=H).status_code == 400            # selection matches nothing
    assert client.post("/api/bta/nope/download-selected", json={"rows": [0]}, headers=H).status_code == 404


def test_download_all_and_current_filter(client, report):
    rid, _ = report
    full = client.get(f"/api/bta/{rid}/download")
    assert len(ids(load_workbook(io.BytesIO(full.data))["Review"])) == 7
    by_dom = client.get(f"/api/bta/{rid}/download?scope=filtered&view=domains&status=No Base Template")
    assert ids(load_workbook(io.BytesIO(by_dom.data))["Review"]) == ["D5", "D6"]
    by_row = client.get(f"/api/bta/{rid}/download?scope=filtered&view=rows&flag=MISSING_URL")
    assert ids(load_workbook(io.BytesIO(by_row.data))["Review"]) == ["D3", "D6"]


def test_page_has_domain_overview_and_selection_controls(client):
    html = client.get("/base-template-analysis").data.decode()
    for needle in ('id="dlsel"', 'id="selfiltered"', 'id="domcards"', "Download All Data", "By Domain"):
        assert needle in html
