"""Multi-PDF finder: scope selectors, link extraction, the page/browser decision, the HTTP flow, and the analysis hand-off."""
import io
import time

import pytest
from flask import Flask
from openpyxl import load_workbook

from tests.helpers import ENGLISH, PDF_HEADERS, FakeResp, make_analyzer, text_pdf
from url_analyzer.multipdf import browser as browser_mod
from url_analyzer.multipdf import scope as S
from url_analyzer.multipdf.extract import count_expanders, extract
from url_analyzer.multipdf.finder import FindError, PdfFinder
from url_analyzer.ui.url_analysis import create_url_analysis_blueprint
from url_analyzer.services.url_analysis_service import UrlAnalysisService

HTML = {"Content-Type": "text/html; charset=utf-8"}
PAGE = "https://example.gov/code"
USER_SCOPE = "#cmctree>.tlist>Li:Nth-of-type(1) a[data-pdf*='.pdf']"       # the selector from the request, capital letters and all


def tree_html(*, expanded=True):
    """A replica of the structure of the real page (titles > divisions > sections), already expanded."""
    leaf = lambda n, folder: (f'<li class="tleaf"><span class="tsec">&sect;</span><a class="tlabel tseclink" href="/code/{n}" data-id="{n}" '
                              f'data-title="Section {n}" data-pdf="/docs/{folder}/{n}.pdf">&sect; {n}. Short title</a></li>')
    t1 = ('<li class="tnode" data-id="1"><div class="trow"><button class="ttoggle" aria-expanded="true">&#9656;</button><span class="tico">&#128193;</span>'
          '<a class="tlabel" href="/code/1">Title 1: Government</a></div><ul class="tchildren">'
          '<li class="tnode"><div class="trow"><button class="ttoggle" aria-expanded="true">&#9656;</button><a class="tlabel" href="/code/2">General Provisions</a></div>'
          f'<ul class="tchildren">{leaf(101, "T1")}{leaf(102, "T1")}</ul></li>'
          f'<li class="tnode"><div class="trow"><a class="tlabel" href="/code/3">Division 2</a></div><ul class="tchildren">{leaf(201, "T1")}</ul></li>'
          '</ul></li>')
    t2 = ('<li class="tnode" data-id="9"><div class="trow"><button class="ttoggle" aria-expanded="true">&#9656;</button>'
          f'<a class="tlabel" href="/code/9">Title 2: Natural Resources</a></div><ul class="tchildren">{leaf(901, "T2")}</ul></li>')
    return f'<html><body><div id="cmctree"><ul class="tlist">{t1}{t2}</ul></div></body></html>'


# ── scope selectors ──────────────────────────────────────────────────────────
def test_split_the_requested_scope():
    root, last = S.split_last_compound(USER_SCOPE)
    assert root == "#cmctree>.tlist>Li:Nth-of-type(1)" and last == "a[data-pdf*='.pdf']"
    assert S.roots(USER_SCOPE) == ["#cmctree>.tlist>Li:Nth-of-type(1)"]
    assert S.attributes(USER_SCOPE) == ["data-pdf"]


@pytest.mark.parametrize("sel,root,last", [
    ("a.pdf", "", "a.pdf"),
    ("#list a.pdf", "#list", "a.pdf"),
    (".year > ul   >   li a", ".year > ul   >   li", "a"),   # spacing is kept: it is equivalent CSS
    ("div + a[title='a b']", "div", "a[title='a b']"),
    ("ul:not(.x .y) a", "ul:not(.x .y)", "a"),
    ("a[href*='x y z']", "", "a[href*='x y z']"),
])
def test_split_edge_cases(sel, root, last):
    assert S.split_last_compound(sel) == (root, last)


def test_selector_lists_and_whole_page_roots():
    assert S.alternatives("a.x, b.y") == ["a.x", "b.y"]
    assert S.roots("#a a.x, #b a.y") == ["#a", "#b"]
    assert S.roots("#a a.x, a.y") == []                       # one alternative can match anywhere -> the whole page
    assert S.attributes("a[data-pdf], a[data-file$='.pdf']") == ["data-pdf", "data-file"]


def test_invalid_selectors_give_a_friendly_error():
    for bad in ("a[", "a::bogus((", "#"):
        with pytest.raises(S.ScopeError, match="not valid CSS"):
            S.validate(bad)
    with pytest.raises(S.ScopeError, match="too long"):
        S.validate("a" * 700)
    assert S.validate("  ") == "" and S.validate(USER_SCOPE) == USER_SCOPE


# ── extraction ───────────────────────────────────────────────────────────────
def test_scope_takes_only_the_first_title_and_resolves_relative_addresses():
    r = extract(tree_html(), "https://example.gov/code#gsc.tab=0", USER_SCOPE)
    assert [f.url for f in r.found] == [f"https://example.gov/docs/T1/{n}.pdf" for n in (101, 102, 201)]
    assert all(f.attr == "data-pdf" for f in r.found)


def test_titles_outline_path_and_node_ids():
    r = extract(tree_html(), PAGE, USER_SCOPE)
    first = r.found[0]
    assert first.title == "§ 101. Short title" and first.node_id == "101"
    assert first.path == "Title 1: Government > General Provisions"            # icons / arrows are not part of the labels
    assert r.found[2].path == "Title 1: Government > Division 2"


def test_without_a_scope_every_document_link_on_the_page_is_taken():
    r = extract(tree_html(), PAGE)
    assert len(r.found) == 4 and r.found[-1].url.endswith("/docs/T2/901.pdf")


def test_data_attribute_wins_over_an_html_page_href():
    html = '<a href="/page/7" data-pdf="/files/7.pdf">Seven</a>'
    assert extract(html, PAGE).found[0].url == "https://example.gov/files/7.pdf"


def test_container_scope_collects_the_links_inside_it():
    html = '<div id="reports"><a href="/a.pdf">A</a> <a href="/b.html">B</a> <a href="/c.PDF?x=1#p=2">C</a></div><a href="/other.pdf">Other</a>'
    r = extract(html, PAGE, "#reports")
    assert [f.url for f in r.found] == ["https://example.gov/a.pdf", "https://example.gov/c.PDF?x=1"]


def test_a_link_the_scope_selects_is_kept_even_without_a_pdf_ending():
    html = '<ul id="l"><li><a class="dl" href="/download?id=3">Three</a></li><li><a class="x" href="/about">About</a></li></ul>'
    assert [f.url for f in extract(html, PAGE, "#l a.dl").found] == ["https://example.gov/download?id=3"]       # the user chose it
    assert extract(html, PAGE, "#l a.dl", strict=True).found == []                                           # strict: needs .pdf
    assert [f.url for f in extract(html, PAGE, "#l a[href*='download']", strict=True).found] == ["https://example.gov/download?id=3"]


def test_a_container_scope_does_not_trust_every_link_inside_it():
    html = '<div id="c"><a href="/about">About</a><a href="/files/get?id=2" class="pdf-link">Get</a><a href="/x.pdf">X</a></div>'
    assert [f.url for f in extract(html, PAGE, "#c").found] == ["https://example.gov/files/get?id=2", "https://example.gov/x.pdf"]


# a replica of the structure inside the law.go.kr frame: [title link] [HWP icon] [PDF icon] per item, relative extensionless addresses
LAW = ("<html><body><ul class='pconfile'>"
       + "".join(f"<li><a class='blu' href='#AJAX' onclick=\"bylInfoDiv('{n}')\"><img alt='펼침' src='/i/coop.gif'> [별표 {n}] 업무처리기간 {n}</a>"
                 f"<a href='flDownload.do?flSeq={n}1&flNm=h{n}' target='_blank'><img alt='HWP파일 다운로드' src='/i/btn_han_16.png'></a>"
                 f"<a href='flDownload.do?flSeq={n}2&flNm=p{n}' target='_blank'><img alt='PDF파일 다운로드' src='/i/btn_pdf.gif'></a></li>" for n in (1, 2, 3))
       + "</ul><ul class='share'><li><a href='javascript:share()'><img alt='카카오톡' src='k.gif'></a></li></ul></body></html>")
LAW_FRAME = "https://www.law.go.kr/LSW//admRulInfoP.do?admRulSeq=1&chrClsCd=010201"


def test_icon_only_pdf_links_in_a_frame_document_with_the_users_scope():
    r = extract(LAW, LAW_FRAME, "ul li a:nth-of-type(3)")
    assert [f.url for f in r.found] == [f"https://www.law.go.kr/LSW/flDownload.do?flSeq={n}2&flNm=p{n}" for n in (1, 2, 3)]
    assert r.found[0].title == "[별표 1] 업무처리기간 1"                  # named by its list item, in Korean, not by the icon's alt text
    assert extract(LAW, LAW_FRAME, "ul li a:nth-of-type(3)", strict=True).found == []


def test_without_a_scope_a_link_that_says_it_is_a_pdf_is_recognised():
    r = extract(LAW, LAW_FRAME)
    assert len(r.found) == 3 and all("flSeq=" in f.url and f.url.endswith(f"p{i + 1}") for i, f in enumerate(r.found))   # the HWP icons are not taken


def test_non_latin_text_counts_as_a_title():
    html = '<ul><li><a href="/a.pdf">제1조 목적</a></li><li><a href="/b.pdf">第一章 总则</a></li><li><a href="/c.pdf">مادة ١</a></li></ul>'
    assert [f.title for f in extract(html, PAGE).found] == ["제1조 목적", "第一章 总则", "مادة ١"]


def test_a_viewer_frame_yields_the_document_inside_it():
    html = '<iframe src="/pdfjs/web/viewer.html?file=%2Fdocs%2Fa.pdf"></iframe>'
    assert [f.url for f in extract(html, PAGE).found] == ["https://example.gov/docs/a.pdf"]


def test_word_files_only_when_asked():
    html = '<a href="/a.pdf">A</a><a href="/b.docx">B</a><a href="/c.doc">C</a>'
    assert len(extract(html, PAGE).found) == 1
    assert len(extract(html, PAGE, types=("pdf", "word")).found) == 3


def test_duplicates_fragments_base_href_and_onclick():
    html = ('<base href="https://cdn.example.org/files/"><a href="a.pdf#page=1">A</a><a href="a.pdf#page=9">again</a>'
            '<button onclick="window.open(\'/b.pdf\')">B</button><a href="mailto:x@y.z">m</a><a href="javascript:void(0)">j</a>')
    r = extract(html, PAGE)
    assert [f.url for f in r.found] == ["https://cdn.example.org/files/a.pdf", "https://cdn.example.org/b.pdf"] and r.duplicates == 1


def test_maximum_is_enforced_and_reported():
    html = "".join(f'<a href="/d/{i}.pdf">{i}</a>' for i in range(10))
    r = extract(html, PAGE, max_items=4)
    assert len(r.found) == 4 and r.truncated


def test_title_falls_back_to_title_attribute_and_aria_label():
    html = '<a href="/a.pdf" title="Annual report"><img src="x.png"></a><a href="/b.pdf" aria-label="Budget"></a>'
    assert [f.title for f in extract(html, PAGE).found] == ["Annual report", "Budget"]


def test_count_expanders_only_inside_the_scopes_region():
    html = tree_html().replace('aria-expanded="true"', 'aria-expanded="false"')
    assert count_expanders(html) == 3
    assert count_expanders(html, USER_SCOPE) == 2                      # the buttons of Title 1 and its first division only
    assert count_expanders("<details><summary>Open</summary><a href='/a.pdf'>x</a></details>") == 1
    assert count_expanders("<div class='opener'></div>", "", ".opener") == 1


# ── the finder: page or browser? ─────────────────────────────────────────────
def pdf_resp(pages=2):
    return FakeResp(200, text_pdf(ENGLISH, pages), {**PDF_HEADERS, "Content-Length": "2048"})


class FakeBrowser:
    """Stands in for the real browser step; records how it was called."""
    def __init__(self, snapshots, clicks=5, notes=None, error=None):
        self.snapshots, self.clicks, self.notes, self.error, self.calls = snapshots, clicks, notes or [], error, []

    def __call__(self, url, **kw):
        self.calls.append((url, kw))
        if self.error:
            raise self.error
        if kw.get("progress"):
            kw["progress"]({"clicks": self.clicks, "links": 3, "message": "Opening sections"})
        return browser_mod.BrowserResult(snapshots=self.snapshots, clicks=self.clicks, rounds=2, pages=len(self.snapshots), notes=self.notes)


def finder(routes, browser=None, available=True):
    a, sess = make_analyzer(routes)
    return PdfFinder(a, browser_fn=browser or FakeBrowser([]), browser_ok=lambda: available), sess


def collapsed_page():
    return page_html('<div id="cmctree"><ul class="tlist"><li class="tnode"><button class="ttoggle" aria-expanded="false">+</button>'
                     '<a class="tlabel" href="/code/1">Title 1</a><ul class="tchildren" hidden></ul></li></ul></div>')


def page_html(body):
    return FakeResp(200, f"<html><body>{body}</body></html>".encode(), HTML)


def test_plain_page_is_read_without_a_browser():
    br = FakeBrowser([])
    f, _ = finder({PAGE: page_html('<a href="/a.pdf">A</a><a href="/b.pdf">B</a>')}, br)
    d = f.find(PAGE)
    assert d.mode == "page" and [x.url for x in d.found] == ["https://example.gov/a.pdf", "https://example.gov/b.pdf"] and br.calls == []


def test_collapsed_sections_trigger_the_browser_and_its_result_is_used():
    br = FakeBrowser([(PAGE, tree_html())], clicks=7)
    f, _ = finder({PAGE: collapsed_page()}, br)
    progress = []
    d = f.find(PAGE + "#gsc.tab=0", USER_SCOPE, progress=progress.append)
    assert d.mode == "browser" and len(d.found) == 3 and d.stats["clicks"] == 7 and d.stats["collapsed_sections"] == 1
    assert d.notes[0].startswith("Opened 7 collapsed sections")
    url, kw = br.calls[0]
    assert kw["scope"] == USER_SCOPE and callable(kw["validate"]) and "user_agent" in kw
    assert any(p.get("clicks") == 7 for p in progress)


def test_pages_of_results_are_merged_and_deduplicated():
    p1 = '<a href="/a.pdf">A</a><a href="/b.pdf">B</a>'
    p2 = '<a href="/b.pdf">B</a><a href="/c.pdf">C</a>'
    br = FakeBrowser([(PAGE, f"<html>{p1}</html>"), (PAGE + "?page=2", f"<html>{p2}</html>")])
    f, _ = finder({PAGE: page_html("<p>list</p>")}, br)
    d = f.find(PAGE, next_selector="a.next")
    assert [x.url.rsplit("/", 1)[1] for x in d.found] == ["a.pdf", "b.pdf", "c.pdf"] and d.stats["duplicates"] == 1 and d.stats["pages"] == 2
    assert br.calls[0][1]["next_selector"] == "a.next"


def test_page_with_nothing_static_falls_back_to_the_browser():
    br = FakeBrowser([(PAGE, '<html><a href="/js.pdf">J</a></html>')])
    f, _ = finder({PAGE: page_html("<div id=app></div>")}, br)
    assert [x.url for x in f.find(PAGE).found] == ["https://example.gov/js.pdf"] and len(br.calls) == 1


def test_mode_page_never_uses_the_browser_and_mode_browser_always_does():
    br = FakeBrowser([(PAGE, "<html><a href='/z.pdf'>Z</a></html>")])
    f, _ = finder({PAGE: collapsed_page()}, br)
    assert f.find(PAGE, mode="page").found == [] and br.calls == []
    f2, _ = finder({PAGE: page_html('<a href="/a.pdf">A</a>')}, br)
    d = f2.find(PAGE, mode="browser")
    assert d.mode == "browser" and len(br.calls) == 1 and "https://example.gov/z.pdf" in [x.url for x in d.found]


def test_without_playwright_the_visible_links_are_still_returned_with_a_note():
    html = '<a href="/a.pdf">A</a><div><button aria-expanded="false">more</button></div>'
    f, _ = finder({PAGE: page_html(html)}, available=False)
    d = f.find(PAGE)
    assert [x.url for x in d.found] == ["https://example.gov/a.pdf"] and any("Playwright is not installed" in n for n in d.notes)


def test_without_playwright_and_nothing_visible_is_an_error_that_says_how_to_fix_it():
    f, _ = finder({PAGE: collapsed_page()}, available=False)
    with pytest.raises(FindError, match="pip install playwright"):
        f.find(PAGE, USER_SCOPE)
    with pytest.raises(FindError, match="pip install playwright"):
        f.find(PAGE, mode="browser")


def test_browser_failure_does_not_lose_what_the_page_already_showed():
    err = browser_mod.BrowserError("The page could not be opened in the browser: timeout")
    f, _ = finder({PAGE: page_html('<a href="/a.pdf">A</a><button aria-expanded="false">x</button>')}, FakeBrowser([], error=err))
    d = f.find(PAGE)
    assert len(d.found) == 1 and any("could not be opened" in n for n in d.notes)
    f2, _ = finder({PAGE: collapsed_page()}, FakeBrowser([], error=err))
    with pytest.raises(FindError, match="could not be opened"):
        f2.find(PAGE)


def test_the_address_itself_being_a_pdf_is_analysed_directly():
    f, _ = finder({"https://example.gov/a.pdf": pdf_resp()})
    d = f.find("https://example.gov/a.pdf")
    assert [x.url for x in d.found] == ["https://example.gov/a.pdf"] and any("itself is a document" in n for n in d.notes)


def test_bad_input_is_refused_before_anything_is_fetched():
    f, sess = finder({})
    for args, msg in ((("",), "Enter the address"), (("http://127.0.0.1/x",), "private or local"), (("ftp://example.gov/x",), "Unsupported URL scheme"),
                      ((PAGE, "a["), "not valid CSS"), ((PAGE, "a", "b[", ""), "click-to-open"), ((PAGE, "a", "", "c(("), "next-page")):
        with pytest.raises(FindError, match=msg):
            f.check(*args)
    with pytest.raises(FindError, match="Unknown mode"):
        f.find(PAGE, mode="telepathy")
    assert sess.calls == []


def test_http_errors_are_reported_plainly():
    f, _ = finder({PAGE: FakeResp(404, b"", HTML)})
    with pytest.raises(FindError, match="HTTP 404"):
        f.find(PAGE)


def test_a_cloudflare_beacon_is_not_reported_as_a_verification_page():
    page = page_html('<p>No links here</p><script>s.src="/cdn-cgi/challenge-platform/scripts/jsd/main.js"</script>')
    f, _ = finder({PAGE: page}, FakeBrowser([(PAGE, "<html></html>")]))
    d = f.find(PAGE)
    assert d.found == [] and not any("bot-protection" in n for n in d.notes)


def test_maximum_number_of_pdfs():
    body = "".join(f'<a href="/d/{i}.pdf">{i}</a>' for i in range(30))
    f, _ = finder({PAGE: page_html(body)})
    d = f.find(PAGE, max_items=10)
    assert len(d.found) == 10 and any("maximum of 10" in n for n in d.notes)


# ── HTTP flow: search -> hand-off to the normal analysis -> results and Excel ─
def make_app(tmp_path, browser):
    pdfs = {f"https://example.gov/docs/T1/{n}.pdf": pdf_resp(n % 5 + 1) for n in (101, 102, 201)}
    a, _ = make_analyzer({PAGE: collapsed_page(), **pdfs})
    fnd = PdfFinder(a, browser_fn=browser, browser_ok=lambda: True)
    svc = UrlAnalysisService(tmp_path / "out", a)
    app = Flask(__name__, template_folder="../templates")
    app.register_blueprint(create_url_analysis_blueprint(tmp_path / "out", tmp_path / "up", service=svc, finder=fnd))
    app.testing = True
    return app.test_client()


@pytest.fixture
def app_env(tmp_path):
    br = FakeBrowser([(PAGE, tree_html())], clicks=3)
    return type("E", (), {"client": make_app(tmp_path, br), "browser": br})


H = {"X-Requested-With": "fetch"}


def wait(client, url, until=lambda s: s["state"] != "running"):
    for _ in range(300):
        s = client.get(url).get_json()
        if until(s):
            return s
        time.sleep(0.02)
    raise AssertionError("timeout")


def test_search_then_analysis_end_to_end(app_env):
    c = app_env.client
    r = c.post("/api/url-analysis/find-pdfs", json={"url": PAGE + "#gsc.tab=0", "scope": USER_SCOPE}, headers=H)
    assert r.status_code == 200
    st = wait(c, f"/api/url-analysis/find-pdfs/{r.get_json()['find_id']}")
    assert st["state"] == "done" and st["found"] == 3 and st["mode_used"] == "browser" and st["clicks"] == 3 and st["job_id"]
    assert any("Opened 3 collapsed sections" in n for n in st["notes"])
    job = st["job_id"]
    status, events = None, {}
    for _ in range(300):
        s = c.get(f"/api/url-analysis/{job}/status?from={len(events)}").get_json()
        events.update({e["idx"]: e for e in s["events"]})
        if s["state"] != "running":
            status = s
            break
        time.sleep(0.02)
    assert status["source"] == "Multi-PDF Page" and status["source_name"] == "example.gov" and len(events) == 3
    row = events[0]
    assert (row["content_type"], row["status"], row["page_count"], row["last_modified"]) == ("PDF", "WORKING", "2", "2025-10-01")
    assert row["book_title"].startswith("§ 101") and row["extra"]["Section Path"] == "Title 1: Government > General Provisions"
    assert row["extra"]["Scope"] == USER_SCOPE and row["extra"]["Source Page"].startswith(PAGE)
    # the shared exporter writes the extra columns, for everything and for a selection
    for url, post in ((f"/api/url-analysis/{job}/download", None), (f"/api/url-analysis/{job}/download", {"indices": [0, 2]})):
        resp = c.post(url, json=post) if post else c.get(url)
        ws = load_workbook(io.BytesIO(resp.data))["URL Analysis"]
        head = [x.value for x in ws[1]]
        assert {"Section Path", "Source Page", "Page Count", "Last Modified", "URL"} <= set(head)
        assert len(list(ws.iter_rows(min_row=2))) == (3 if post is None else 2)


def test_search_input_validation_and_guards(app_env):
    c = app_env.client
    assert c.post("/api/url-analysis/find-pdfs", json={"url": PAGE}).status_code == 400                       # missing header
    for body, msg in (({"url": ""}, "Enter the address"), ({"url": "http://127.0.0.1/x"}, "private"), ({"url": PAGE, "scope": "a["}, "not valid CSS"),
                      ({"url": PAGE, "mode": "x"}, "Unknown mode"), ({"url": PAGE, "max_items": "many"}, "must be a number")):
        r = c.post("/api/url-analysis/find-pdfs", json=body, headers=H)
        assert r.status_code == 400 and msg in r.get_json()["error"], body
    assert c.get("/api/url-analysis/find-pdfs/nope").status_code == 404
    assert c.post("/api/url-analysis/find-pdfs/nope/cancel", headers=H).status_code == 404
    assert c.post("/api/url-analysis/find-pdfs/nope/cancel").status_code == 400
    assert c.get("/api/url-analysis/find-pdfs/capabilities").get_json() == {"browser": True}


def test_no_pdfs_found_is_a_clear_error_with_a_hint(app_env):
    c = app_env.client
    app_env.browser.snapshots = [(PAGE, "<html><body><a class='x' href='/page.html'>p</a></body></html>")]
    fid = c.post("/api/url-analysis/find-pdfs", json={"url": PAGE, "scope": "a.nothing"}, headers=H).get_json()["find_id"]
    st = wait(c, f"/api/url-analysis/find-pdfs/{fid}")
    assert st["state"] == "error" and "No PDF links were found" in st["error"] and "scope" in st["error"]
    assert st["job_id"] is None


def test_search_can_be_stopped(tmp_path):
    class Slow(FakeBrowser):
        def __call__(self, url, **kw):
            for _ in range(500):                                   # waits until the user presses Stop
                if kw["cancelled"]():
                    break
                time.sleep(0.01)
            return browser_mod.BrowserResult(snapshots=[(PAGE, tree_html())], clicks=1, rounds=1, pages=1)

    c = make_app(tmp_path, Slow([]))
    fid = c.post("/api/url-analysis/find-pdfs", json={"url": PAGE}, headers=H).get_json()["find_id"]
    time.sleep(0.2)
    assert c.post(f"/api/url-analysis/find-pdfs/{fid}/cancel", headers=H).get_json() == {"cancelled": True}
    st = wait(c, f"/api/url-analysis/find-pdfs/{fid}")
    assert st["state"] == "cancelled" and st["job_id"] is None


def test_the_page_has_the_new_input_tab_and_form():
    import importlib
    html = importlib.import_module("dashboard").create_app().test_client().get("/url-analysis").data.decode()
    for needle in ('data-t="find"', 'id="p-find"', 'id="f-url"', 'id="f-scope"', 'id="go-find"', "Find PDFs on a Page", "Section Path"):
        assert needle in html


# ── frames: the links are inside an <iframe> document ────────────────────────
FRAME_URL = "https://example.gov/frame/content"


def shell(src="/frame/content"):
    return page_html(f'<iframe id="lawService" name="lawService" src="{src}" style="width:100%;height:100%"></iframe>')


def test_links_inside_a_server_rendered_frame_are_found_without_a_browser():
    br = FakeBrowser([])
    routes = {PAGE: shell(), FRAME_URL: FakeResp(200, LAW.encode("utf-8"), HTML)}
    f, _ = finder(routes, br)
    d = f.find(PAGE, "ul li a:nth-of-type(3)")
    assert d.mode == "page" and len(d.found) == 3 and br.calls == []
    assert d.found[0].url == "https://example.gov/frame/flDownload.do?flSeq=12&flNm=p1" and d.found[0].title.startswith("[별표 1]")


def test_frame_in_a_frame_and_no_scope():
    inner = page_html('<iframe src="/frame/deep"></iframe>')
    routes = {PAGE: shell(), FRAME_URL: inner, "https://example.gov/frame/deep": FakeResp(200, LAW.encode("utf-8"), HTML)}
    f, _ = finder(routes)
    assert len(f.find(PAGE).found) == 3


def test_frame_that_fails_or_is_a_document_does_not_break_the_search():
    routes = {PAGE: page_html('<iframe src="/gone"></iframe><iframe src="/doc"></iframe><a href="/ok.pdf">ok</a>'),
              "https://example.gov/gone": FakeResp(500, b"", HTML), "https://example.gov/doc": pdf_resp()}
    d = finder(routes)[0].find(PAGE)
    assert sorted(x.url.rsplit("/", 1)[1] for x in d.found) == ["doc", "ok.pdf"]


def test_collapsed_sections_inside_a_frame_trigger_the_browser():
    inner = FakeResp(200, b'<html><body><div id="t"><button aria-expanded="false">+</button></div></body></html>', HTML)
    br = FakeBrowser([(FRAME_URL, LAW)])
    d = finder({PAGE: shell(), FRAME_URL: inner}, br)[0].find(PAGE, "ul li a:nth-of-type(3)")
    assert d.mode == "browser" and len(d.found) == 3 and d.stats["collapsed_sections"] == 1


def test_a_frame_page_with_nothing_static_uses_the_browser_snapshots_of_the_frame():
    inner = FakeResp(200, b"<html><body><ul id='empty'></ul></body></html>", HTML)           # filled by JavaScript in real life
    br = FakeBrowser([(PAGE, '<html><iframe src="/frame/content"></iframe></html>'), (FRAME_URL, LAW)])
    d = finder({PAGE: shell(), FRAME_URL: inner}, br)[0].find(PAGE, "ul li a:nth-of-type(3)")
    assert d.mode == "browser" and len(d.found) == 3
