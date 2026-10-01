"""The original analysis spec, run by the shared engine for every input path:
status classification, direct documents, landing pages (SINGLE / MULTIPLE / NONE), link verification,
de-duplication, retries / back-off / 429, browser fallback hook, SUMMARY sheet."""
import io
import socket
import threading
import time

import pandas as pd
import pytest
import requests
from openpyxl import load_workbook

from tests.helpers import ENGLISH, PDF_HEADERS, FakeResp, make_analyzer, text_pdf
from url_analyzer.analysis import BookRecord, analyze_urls, export_analysis_results, summarize

HTML = {"Content-Type": "text/html; charset=utf-8"}
PAGE = "https://example.gov/report"


def page(body: str) -> FakeResp:
    return FakeResp(200, f"<html><body>{body}</body></html>".encode(), HTML)


def pdf(pages=5, text=ENGLISH, **extra) -> FakeResp:
    return FakeResp(200, text_pdf(text, pages), {**PDF_HEADERS, **extra})


# ── direct documents ─────────────────────────────────────────────────────────
def test_pdf_detected_by_content_type_even_without_an_extension():
    a, _ = make_analyzer({"https://example.gov/download?id=1": pdf(2)})
    r = a.analyze("https://example.gov/download?id=1")
    assert (r.url_type, r.document_link_type, r.page_count, r.content_type) == ("DIRECT_PDF", "DIRECT", "2", "PDF")
    assert r.document_url == "https://example.gov/download?id=1" and r.document_link_count is None


def test_redirect_to_a_pdf_analyses_the_final_url():
    a, _ = make_analyzer({"https://example.gov/document": FakeResp(301, b"", {"Location": "/document.pdf"}),
                          "https://example.gov/document.pdf": pdf(7)})
    r = a.analyze("https://example.gov/document")
    assert (r.url_type, r.redirects, r.http_status, r.final_status, r.page_count) == ("DIRECT_PDF", 1, 301, 200, "7")
    assert r.final_url == r.document_url == "https://example.gov/document.pdf"


def test_scanned_pdf_has_no_text_and_unknown_language():
    a, _ = make_analyzer({"https://example.gov/scan.pdf": pdf(3, text="")})
    r = a.analyze("https://example.gov/scan.pdf")
    assert (r.text_extraction_status, r.language, r.page_count) == ("NO_TEXT", "Unknown", "3")


def test_direct_document_type_for_doc():
    a, _ = make_analyzer({"https://example.gov/a.doc": FakeResp(200, b"\xd0\xcf\x11\xe0" + b"x" * 20, {"Content-Type": "application/msword"})})
    r = a.analyze("https://example.gov/a.doc")
    assert (r.url_type, r.document_link_type, r.direct_document) == ("DIRECT_DOCUMENT", "DIRECT", True)


def test_non_html_non_document_is_unknown():
    a, _ = make_analyzer({"https://example.gov/pic": FakeResp(200, b"\x89PNG....", {"Content-Type": "image/png"})})
    r = a.analyze("https://example.gov/pic")
    assert (r.status, r.url_type, r.document_link_type) == ("WORKING", "UNKNOWN", "")


# ── landing pages ────────────────────────────────────────────────────────────
def test_landing_page_without_documents_is_NONE():
    a, _ = make_analyzer({PAGE: page('<a href="/about">About</a><a href="/news.html">News</a>')})
    r = a.analyze(PAGE)
    assert (r.status, r.url_type, r.document_link_type, r.document_link_count, r.document_url) == ("WORKING", "LANDING_PAGE", "NONE", 0, "")
    assert r.page_count == "N/A" and r.analysis_status == "Success"


def test_single_document_link_is_followed_and_its_metadata_returned():
    doc = "https://example.gov/files/report.pdf"
    a, sess = make_analyzer({PAGE: page('<a href="/files/report.pdf">Download PDF</a>'), doc: pdf(35)})
    r = a.analyze(PAGE)
    assert (r.url_type, r.document_link_type, r.document_link_count) == ("LANDING_PAGE", "SINGLE", 1)
    assert r.document_url == doc and r.page_count == "35" and r.language == "English"
    assert r.last_modified == "2025-10-01" and r.text_extraction_status == "OK" and r.file_size_bytes
    assert not r.direct_document and r.content_type == "HTML" and r.analysis_status == "Success"
    assert [c[1] for c in sess.calls] == [PAGE, doc]


def test_links_are_normalised_and_deduplicated_before_counting():
    doc = "https://example.gov/files/report.pdf"
    a, sess = make_analyzer({PAGE: page(
        '<a href="/files/report.pdf">Download PDF</a><a href="files/report.pdf#page=2">View PDF</a>'
        f'<a href="{doc}">PDF version</a><iframe src="/files/report.pdf"></iframe>'), doc: pdf(4)})
    r = a.analyze(PAGE)
    assert r.document_link_type == "SINGLE" and r.document_link_count == 1 and r.page_count == "4"
    assert [c[1] for c in sess.calls].count(doc) == 1                         # the document is fetched once


def test_multiple_document_links_are_reported_not_chosen():
    names = ["ch1.pdf", "ch2.pdf", "ch3.pdf", "appendix.docx"]
    a, sess = make_analyzer({PAGE: page("".join(f'<a href="/{n}">{n}</a>' for n in names))})
    r = a.analyze(PAGE)
    assert (r.document_link_type, r.document_link_count) == ("MULTIPLE", 4)
    assert r.document_urls == [f"https://example.gov/{n}" for n in names]
    assert r.page_count == "N/A" and r.document_url == "" and len(sess.calls) == 1   # nothing was followed


def test_iframe_embed_and_object_links_count():
    a, _ = make_analyzer({PAGE: page('<iframe src="/a.pdf"></iframe><embed src="/b.pdf"><object data="/c.docx"></object>')})
    assert a.analyze(PAGE).document_link_count == 3


def test_download_style_link_is_verified_by_content_type():
    a, sess = make_analyzer({PAGE: page('<a href="/download?id=7">Download</a>'), "https://example.gov/download?id=7": pdf(9)})
    r = a.analyze(PAGE)
    assert (r.document_link_type, r.document_url, r.page_count) == ("SINGLE", "https://example.gov/download?id=7", "9")


def test_anchor_text_alone_does_not_make_a_document():
    a, _ = make_analyzer({PAGE: page('<a href="/about-us">Download PDF</a>'), "https://example.gov/about-us": page("About us")})
    r = a.analyze(PAGE)
    assert r.document_link_type == "NONE" and r.document_link_count == 0


def test_unverifiable_candidates_are_ignored():
    a, _ = make_analyzer({PAGE: page('<a href="/download?id=1">Download</a>'), "https://example.gov/download?id=1": FakeResp(404, b"", HTML)})
    assert a.analyze(PAGE).document_link_type == "NONE"


def test_single_link_that_is_broken_is_reported_clearly():
    a, _ = make_analyzer({PAGE: page('<a href="/gone.pdf">x</a>'), "https://example.gov/gone.pdf": FakeResp(404, b"", HTML)})
    r = a.analyze(PAGE)
    assert (r.status, r.document_link_type, r.analysis_status, r.page_count) == ("WORKING", "SINGLE", "Unknown", "Unknown")
    assert r.error == "Linked document could not be analysed: HTTP 404 - Not found"


def test_single_link_that_returns_html_is_not_treated_as_a_document():
    a, _ = make_analyzer({PAGE: page('<a href="/fake.pdf">x</a>'), "https://example.gov/fake.pdf": page("oops")})
    r = a.analyze(PAGE)
    assert r.analysis_status == "Unknown" and "could not be analysed" in r.error and r.document_link_type == "SINGLE"


def test_followed_document_cannot_trigger_further_crawling():
    # the linked "document" is itself a page with links: it is NOT followed again
    a, sess = make_analyzer({PAGE: page('<a href="/one.pdf">x</a>'), "https://example.gov/one.pdf": page('<a href="/two.pdf">y</a>')})
    a.analyze(PAGE)
    assert [c[1] for c in sess.calls] == [PAGE, "https://example.gov/one.pdf"]


# ── browser fallback hook ────────────────────────────────────────────────────
class FakeRenderer:
    def __init__(self, html):
        self.html, self.calls = html, []

    def render(self, url):
        self.calls.append(url)
        return self.html


def test_browser_fallback_finds_javascript_generated_links():
    doc = "https://example.gov/js.pdf"
    rend = FakeRenderer('<a href="/js.pdf">PDF</a>')
    a, _ = make_analyzer({PAGE: page("<div id=app></div>"), doc: pdf(6)})
    a.renderer = rend
    r = a.analyze(PAGE)
    assert (r.document_link_type, r.page_count, rend.calls) == ("SINGLE", "6", [PAGE])
    assert any("browser" in n for n in r.notes)


def test_browser_fallback_is_not_used_when_static_html_has_links_or_when_disabled():
    rend = FakeRenderer("")
    a, _ = make_analyzer({PAGE: page('<a href="/a.pdf">a</a><a href="/b.pdf">b</a>')})
    a.renderer = rend
    assert a.analyze(PAGE).document_link_type == "MULTIPLE" and rend.calls == []
    b, _ = make_analyzer({PAGE: page("nothing")})
    assert b.renderer is None and b.analyze(PAGE).document_link_type == "NONE"


# ── status classification, retries, back-off, rate limiting ──────────────────
@pytest.mark.parametrize("code,state", [(200, "WORKING"), (404, "NOT_WORKING"), (410, "NOT_WORKING"), (500, "NOT_WORKING"),
                                        (502, "NOT_WORKING"), (503, "NOT_WORKING"), (504, "NOT_WORKING"),
                                        (403, "BLOCKED"), (401, "BLOCKED")])
def test_status_classification(code, state):
    a, _ = make_analyzer({PAGE: FakeResp(code, b"<html></html>", HTML)})
    assert a.analyze(PAGE).status == state


def test_timeout_dns_connection_and_invalid_urls():
    def gaierror(host):
        raise socket.gaierror("x")
    a, _ = make_analyzer({"https://slow.example/a": requests.exceptions.ReadTimeout("t"),
                          "https://down.example/a": requests.exceptions.ConnectionError("refused"),
                          "https://ssl.example/a": requests.exceptions.SSLError("bad")})
    assert a.analyze("https://slow.example/a").status == "TIMEOUT"
    assert a.analyze("https://down.example/a").status == "ERROR"
    assert a.analyze("https://ssl.example/a").status == "ERROR"
    assert a.analyze("not a url").status == "ERROR" and a.analyze("not a url").url_type == "NOT_WORKING"
    dns, _ = make_analyzer({}, resolver=gaierror)
    assert dns.analyze("https://nohost.example/").status == "ERROR"


def test_http_429_is_retried_respecting_retry_after_then_succeeds():
    attempts = {"n": 0}

    def route(kw):
        attempts["n"] += 1
        return FakeResp(429, b"", {"Retry-After": "0"}) if attempts["n"] == 1 else page("ok")
    a, _ = make_analyzer({PAGE: route}, retries=2, retry_delay=0.01)
    r = a.analyze(PAGE)
    assert r.status == "WORKING" and attempts["n"] == 2


def test_persistent_429_ends_as_blocked():
    a, sess = make_analyzer({PAGE: FakeResp(429, b"", {"Retry-After": "0"})}, retries=1, retry_delay=0.01)
    r = a.analyze(PAGE)
    assert (r.status, r.error, len(sess.calls)) == ("BLOCKED", "HTTP 429 - Rate limited", 2)


def test_exponential_backoff_between_retries():
    a, sess = make_analyzer({PAGE: requests.exceptions.ConnectionError("down")}, retries=2, retry_delay=0.05)
    t0 = time.monotonic()
    r = a.analyze(PAGE)
    assert len(sess.calls) == 3 and time.monotonic() - t0 >= 0.05 + 0.10 - 0.01 and r.status == "ERROR"


def test_request_delay_spaces_requests_to_the_same_host():
    a, _ = make_analyzer({"https://example.gov/1": page("1"), "https://example.gov/2": page("2")}, request_delay=0.15)
    t0 = time.monotonic()
    a.analyze("https://example.gov/1"); a.analyze("https://example.gov/2")
    assert time.monotonic() - t0 >= 0.14


def test_batch_of_landing_pages_on_one_host_does_not_deadlock():
    routes = {f"https://same.example/p{i}": page(f'<a href="/d{i}.pdf">d</a>') for i in range(6)}
    routes.update({f"https://same.example/d{i}.pdf": pdf(2) for i in range(6)})
    a, _ = make_analyzer(routes, per_host_limit=1, max_workers=6)
    out = []
    t = threading.Thread(target=lambda: out.extend(analyze_urls([{"url": f"https://same.example/p{i}"} for i in range(6)], analyzer=a)))
    t.start(); t.join(15)
    assert not t.is_alive(), "deadlock"
    assert [r.document_link_type for r in out] == ["SINGLE"] * 6 and all(r.page_count == "2" for r in out)


# ── SUMMARY sheet and shared export ──────────────────────────────────────────
def mixed_results():
    routes = {"https://a.example/doc.pdf": pdf(10), "https://b.example/one": page('<a href="/x.pdf">x</a>'),
              "https://b.example/x.pdf": pdf(5), "https://c.example/many": page('<a href="/1.pdf">1</a><a href="/2.pdf">2</a>'),
              "https://d.example/none": page("nothing"), "https://e.example/gone": FakeResp(404, b"", HTML),
              "https://f.example/img": FakeResp(200, b"GIF89a", {"Content-Type": "image/gif"})}
    a, _ = make_analyzer(routes)
    recs = [BookRecord(url=u, document_id=f"D{i}", book_title=f"B{i}") for i, u in enumerate(
        list(routes)[:1] + ["https://b.example/one", "https://c.example/many", "https://d.example/none", "https://e.example/gone", "https://f.example/img", "bad url"])]
    return analyze_urls(recs, analyzer=a)


def test_summary_categories_match_the_spec():
    got = dict(summarize(mixed_results()))
    assert got["Total URLs"] == 7 and got["Working"] == 5 and got["Not Working"] == 2
    assert (got["Direct PDF"], got["Direct Document"], got["Landing Page"]) == (1, 0, 3)
    assert (got["Single Document Link"], got["Multiple Document Links"], got["No Document Link"], got["Unknown"]) == (1, 1, 1, 1)
    assert got["Total Pages"] == 15


def test_export_has_summary_document_links_and_spec_columns(tmp_path):
    path = export_analysis_results(mixed_results(), tmp_path / "o.xlsx", source="Excel Batch Upload", source_name="in.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["URL Analysis", "Document Links", "SUMMARY"]
    df = pd.read_excel(path, dtype=object)
    for col in ("Status", "URL Type", "Document URL", "Document Link Type", "Document Link Count", "Text Extraction Status"):
        assert col in df.columns
    row = df[df["Document ID"] == "D2"].iloc[0]
    assert (row["URL Type"], row["Document Link Type"], int(row["Document Link Count"])) == ("LANDING_PAGE", "MULTIPLE", 2)
    links = pd.read_excel(path, sheet_name="Document Links")
    assert sorted(links["Document URL"]) == ["https://b.example/x.pdf", "https://c.example/1.pdf", "https://c.example/2.pdf"]
    summary = dict(pd.read_excel(path, sheet_name="SUMMARY").values)
    assert summary["Total URLs"] == 7 and summary["Source"] == "Excel Batch Upload"


# ── download selected (web) ──────────────────────────────────────────────────
def test_download_selected_exports_only_the_chosen_rows(tmp_path):
    from pathlib import Path
    from flask import Flask
    from url_analyzer.services.url_analysis_service import UrlAnalysisService
    from url_analyzer.ui.url_analysis import create_url_analysis_blueprint
    routes = {f"https://example.gov/p{i}": page("x") for i in range(5)}
    a, _ = make_analyzer(routes)
    app = Flask(__name__, template_folder=str(Path(__file__).parents[1] / "templates"))
    app.register_blueprint(create_url_analysis_blueprint(tmp_path / "o", tmp_path / "u", service=UrlAnalysisService(tmp_path / "o", a)))
    c = app.test_client()
    jid = c.post("/api/url-analysis/analyze", json={"mode": "urls", "text": "\n".join(routes)}).get_json()["job_id"]
    for _ in range(100):
        if c.get(f"/api/url-analysis/{jid}/status").get_json()["state"] != "running":
            break
        time.sleep(0.05)
    sel = c.post(f"/api/url-analysis/{jid}/download", json={"indices": [1, 3]})
    assert sel.status_code == 200 and sel.data[:2] == b"PK"
    df = pd.read_excel(io.BytesIO(sel.data))
    assert sorted(df["URL"]) == ["https://example.gov/p1", "https://example.gov/p3"]
    assert len(pd.read_excel(io.BytesIO(c.get(f"/api/url-analysis/{jid}/download").data))) == 5
    assert c.post(f"/api/url-analysis/{jid}/download", json={"indices": []}).status_code == 400
    assert c.post(f"/api/url-analysis/{jid}/download", json={"indices": ["x"]}).status_code == 400
    assert c.post(f"/api/url-analysis/{jid}/download", json={"indices": [99]}).status_code == 409
