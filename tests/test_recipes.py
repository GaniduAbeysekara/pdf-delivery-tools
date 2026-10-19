"""Site recipes inside 'Find PDFs on a Page': recognising the page, ready-made results in the normal dashboard, reading the files in-session."""
import io
import time

import pytest
from flask import Flask
from openpyxl import load_workbook

from tests.helpers import make_analyzer
from tests.test_cbj_browser import FAST, site  # noqa: F401  (the replica site fixture)
from url_analyzer.multipdf.finder import PdfFinder
from url_analyzer.recipes import cbj, registry
from url_analyzer.recipes.registry import attachments_to_results, find_recipe
from url_analyzer.services.url_analysis_service import UrlAnalysisService
from url_analyzer.ui.url_analysis import create_url_analysis_blueprint

H = {"X-Requested-With": "fetch"}
CBJ_URL = "https://www.superfinanciera.gov.co/publicaciones/10115528/circular-basica-juridica-ce-00625/"


def att(**kw):
    base = dict(parte="Parte I. A", titulo="Título I. B", capitulo="Capítulo III. C", name="Anexo 1.xls", url="https://s.gov/api/descargarAnexo/328",
                ext="xls", attachment_id="328", found_in="api, click", verified_by_click=True, app_ids="1/1/3", page_url="https://s.gov/viewer")
    return cbj.Attachment(**{**base, **kw})


# ── recognising the page ─────────────────────────────────────────────────────
@pytest.mark.parametrize("url,yes", [
    (CBJ_URL, True), ("superfinanciera.gov.co/publicaciones/1/circular-basica-juridica-x/", True),
    ("https://www.superfinanciera.gov.co/visualizadorCBJ/inicio", True), ("https://www.superfinanciera.gov.co/publicaciones/9/otra-cosa/", False),
    ("https://example.org/circular-basica-juridica/", False), ("", False)])
def test_the_recipe_matches_only_its_page(url, yes):
    assert (find_recipe(url) is not None) is yes


def test_the_recipe_describes_its_options_for_the_page():
    d = find_recipe(CBJ_URL).describe()
    assert d["name"].startswith("Superfinanciera") and [o["key"] for o in d["options"]] == ["parte", "read_files"]
    assert d["options"][0]["choices"][0] == ["", "All Partes"]


# ── attachments become ordinary analysis results ─────────────────────────────
def test_unread_attachments_are_listed_honestly_not_as_failures():
    r = attachments_to_results([att(), att(name="Guia.docx", ext="docx", attachment_id="9", url="https://s.gov/api/descargarAnexo/9")], CBJ_URL)
    a = r[0]
    assert (a.status, a.url_type, a.content_type, a.direct_document, a.document_link_type) == ("WORKING", "DIRECT_DOCUMENT", "Other", True, "DIRECT")
    assert (a.analysis_status, a.page_count, a.file_size, a.last_modified) == ("Unknown", "Unknown", "Unknown", "Unknown")
    assert "own session" in a.error and "Read the files" in a.error
    assert a.extra["Section Path"] == "Parte I. A > Título I. B > Capítulo III. C" and a.extra["Site IDs"] == "1/1/3" and a.book_title == "Anexo 1.xls"
    assert r[1].content_type == "DOCX"


def test_read_attachments_carry_size_pages_and_date():
    a = attachments_to_results([att(ext="pdf", name="G.pdf", size_bytes=2048, pages="7", language="English", last_modified="2025-10-01")], CBJ_URL)[0]
    assert (a.analysis_status, a.http_status, a.file_size, a.file_size_bytes, a.page_count, a.language, a.last_modified, a.content_type) == (
        "Success", 200, "2.0 KB", 2048, "7", "English", "2025-10-01", "PDF")
    unreadable = attachments_to_results([att(size_bytes=9, pages="Unknown", read_error="could not open")], CBJ_URL)[0]
    assert unreadable.analysis_status == "Unknown" and unreadable.error == "could not open"


def test_a_listed_attachment_with_no_download_request_is_an_error_row():
    r = attachments_to_results([att(url="", verified_by_click=False)], CBJ_URL)[0]
    assert (r.status, r.analysis_status) == ("ERROR", "Failed") and "no download request" in r.error


# ── the dashboard side: a finished job that is not re-fetched ────────────────
@pytest.fixture
def client(tmp_path):
    a, _ = make_analyzer({})
    svc = UrlAnalysisService(tmp_path / "out", a)
    app = Flask(__name__, template_folder="../templates")
    app.register_blueprint(create_url_analysis_blueprint(tmp_path / "out", tmp_path / "up", service=svc, finder=PdfFinder(a, browser_ok=lambda: True)))
    app.testing = True
    return app.test_client(), svc


def test_prepared_results_form_a_finished_job_with_downloads_and_no_reanalysis(client):
    c, svc = client
    results = attachments_to_results([att(), att(name="B.xlsx", ext="xlsx", attachment_id="2", url="https://s.gov/api/descargarAnexo/2")], CBJ_URL)
    job = svc.start_prepared(results, source="Multi-PDF Page", source_name="s.gov", source_columns=["Section Path", "Source Page", "Site IDs"])["job_id"]
    st = c.get(f"/api/url-analysis/{job}/status?from=0").get_json()
    assert st["state"] == "done" and st["total"] == 2 and st["batch_done"] == 2 and [e["book_title"] for e in st["events"]] == ["Anexo 1.xls", "B.xlsx"]
    ws = load_workbook(io.BytesIO(c.get(f"/api/url-analysis/{job}/download").data))["URL Analysis"]
    head = [x.value for x in ws[1]]
    assert {"Section Path", "Source Page", "Site IDs", "URL", "Page Count", "Last Modified"} <= set(head) and ws.max_row == 3
    assert c.post(f"/api/url-analysis/{job}/download", json={"indices": [1]}).status_code == 200
    assert c.post(f"/api/url-analysis/{job}/reanalyze", json={}).status_code == 409                          # nothing to fetch again
    assert c.post(f"/api/url-analysis/{job}/send-to-base-templates", json={}).status_code == 200


def test_the_recipe_route(client):
    c, _ = client
    assert c.get("/api/url-analysis/find-pdfs/recipe?url=" + CBJ_URL).get_json()["recipe"]["key"] == "superfinanciera-cbj"
    assert c.get("/api/url-analysis/find-pdfs/recipe?url=https://example.org/x").get_json() == {"recipe": None}
    assert c.get("/api/url-analysis/find-pdfs/recipe").get_json() == {"recipe": None}


def test_the_tab_has_the_recipe_panel():
    import importlib
    html = importlib.import_module("dashboard").create_app().test_client().get("/url-analysis").data.decode()
    assert 'id="f-recipe"' in html and "refreshRecipe" in html and "recipe: readRecipe()" in html


# ── through the real Find PDFs job, against the replica site (a real browser) ─
pytestmark_browser = pytest.mark.skipif(not __import__("importlib").util.find_spec("playwright"), reason="Playwright is not installed")


@pytest.fixture
def replica_app(site, tmp_path, monkeypatch):  # noqa: F811
    monkeypatch.setattr(registry.CBJ, "matches", lambda url: "127.0.0.1" in url)             # the replica stands in for superfinanciera.gov.co
    monkeypatch.setattr(cbj, "Timing", lambda **kw: FAST)
    real_crawl = cbj.crawl
    monkeypatch.setattr(cbj, "crawl", lambda url, **kw: real_crawl(url, timing=FAST, **kw))
    a, _ = make_analyzer({})
    a.http.validate = lambda url: ""                                                         # the replica is on 127.0.0.1
    svc = UrlAnalysisService(tmp_path / "out", a)
    app = Flask(__name__, template_folder="../templates")
    app.register_blueprint(create_url_analysis_blueprint(tmp_path / "out", tmp_path / "up", service=svc, finder=PdfFinder(a, browser_ok=lambda: True)))
    app.testing = True
    return app.test_client(), site


def wait(c, url, until=lambda s: s["state"] != "running", tries=1500):
    for _ in range(tries):
        s = c.get(url).get_json()
        if until(s):
            return s
        time.sleep(0.2)
    raise AssertionError("timeout")


@pytestmark_browser
def test_find_pdfs_on_a_recognised_page_gives_the_attachments_in_the_dashboard(replica_app):
    c, site = replica_app
    site.hits.clear()
    fid = c.post("/api/url-analysis/find-pdfs", json={"url": site.url, "recipe": {"parte": ""}}, headers=H).get_json()["find_id"]
    st = wait(c, f"/api/url-analysis/find-pdfs/{fid}")
    assert st["state"] == "done" and st["mode_used"] == "recipe" and st["job_id"]
    assert any("Superfinanciera recipe" in n for n in st["notes"])
    job = c.get(f"/api/url-analysis/{st['job_id']}/status?from=0").get_json()
    attachments = [e for e in job["events"] if "descargarAnexo" in e["url"]]
    # the replica's main page and second iframe each hold a document link, listed under every Capitulo (7 x 2) besides the 5 attachments
    assert job["state"] == "done" and len(attachments) == 5 and len(job["events"]) == st["found"] == 19 and job["source"] == "Multi-PDF Page"
    e = next(e for e in job["events"] if e["document_id"] == "330")
    assert e["book_title"] == "Formato numeral.pdf" and e["extra"]["Section Path"].endswith("Capítulo III. C") and e["extra"]["Site IDs"] == "1/1/3"
    assert e["analysis_status"] == "Unknown" and e["page_count"] == "Unknown"                    # not read: honest, not a failure
    assert not any("descargarAnexo" in h for h in site.hits)                                      # nothing was downloaded


@pytestmark_browser
def test_reading_the_files_gives_size_pages_and_date(replica_app):
    c, site = replica_app
    fid = c.post("/api/url-analysis/find-pdfs", json={"url": site.url, "recipe": {"parte": "1", "read_files": True}}, headers=H).get_json()["find_id"]
    st = wait(c, f"/api/url-analysis/find-pdfs/{fid}")
    assert st["state"] == "done"
    events = {e["document_id"] + str(i): e for i, e in enumerate(c.get(f"/api/url-analysis/{st['job_id']}/status?from=0").get_json()["events"]) if "descargarAnexo" in e["url"]}
    assert len(events) == 4                                                                       # Parte I only: 328, 329, 330 and 328 again
    pdf = next(e for e in events.values() if e["document_id"] == "330")
    assert (pdf["analysis_status"], pdf["page_count"], pdf["content_type"], pdf["last_modified"], pdf["http_status"]) == ("Success", "3", "PDF", "2025-10-01", "200")
    assert pdf["file_size_bytes"] > 500 and pdf["language"] == "English"
    xls = next(e for e in events.values() if e["document_id"] == "328")
    assert xls["file_size_bytes"] == 9 and xls["last_modified"] == "2025-10-01" and xls["page_count"] == "Unknown"
    assert any("descargarAnexo" in h for h in site.hits)                                           # this time the files were really requested (in memory only)


@pytestmark_browser
def test_a_recognised_page_can_be_stopped(replica_app):
    c, site = replica_app
    fid = c.post("/api/url-analysis/find-pdfs", json={"url": site.url}, headers=H).get_json()["find_id"]
    wait(c, f"/api/url-analysis/find-pdfs/{fid}", lambda s: s["state"] != "running" or "Reading" in (s["message"] or "") or "Capítulo" in (s["message"] or ""))
    assert c.post(f"/api/url-analysis/find-pdfs/{fid}/cancel", headers=H).get_json() == {"cancelled": True}
    st = wait(c, f"/api/url-analysis/find-pdfs/{fid}")
    assert st["state"] in ("cancelled", "done")


@pytestmark_browser
def test_the_page_mode_option_skips_the_recipe(replica_app):
    c, site = replica_app
    fid = c.post("/api/url-analysis/find-pdfs", json={"url": site.url, "mode": "page"}, headers=H).get_json()["find_id"]
    st = wait(c, f"/api/url-analysis/find-pdfs/{fid}")
    assert st.get("mode_used") != "recipe"
