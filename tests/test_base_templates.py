"""Base Templates page: store rules, Excel import/export, HTTP API, navigation."""
import io
import sys
from pathlib import Path

import pytest
from flask import Flask
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

from url_analyzer.base_templates import excel
from url_analyzer.base_templates.store import BaseTemplateStore, ValidationError, clean, normalize_domains
from url_analyzer.base_templates.web import create_base_templates_blueprint

H = {"X-Requested-With": "fetch"}
ROW = {"created_date": "2026-09-01", "template": "TPL-A", "domain": "a.example.org", "added_to_replit": "Yes",
       "dev_name": "Dev One", "comments": "first"}


@pytest.fixture
def store(tmp_path):
    return BaseTemplateStore(tmp_path / "bt.json")


@pytest.fixture
def client(tmp_path):
    app = Flask(__name__, template_folder=str(Path(__file__).resolve().parent.parent / "templates"))
    app.register_blueprint(create_base_templates_blueprint(tmp_path / "bt.json"))
    app.testing = True
    return app.test_client()


def make_xlsx(rows, sheet="Base Template", title_rows=0, header=None):
    wb = Workbook()
    wb.active.title = "Other"
    ws = wb.create_sheet(sheet)
    for _ in range(title_rows):
        ws.append(["Base templates tracker"])
    ws.append(header or ["Created Date", "Monitoring Templates", "Cached Domain", "Added to Replit?", "Dev Name", "Comments"])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── store ────────────────────────────────────────────────────────────────────
def test_add_persists_and_reloads(store, tmp_path):
    store.add(ROW)
    assert BaseTemplateStore(tmp_path / "bt.json").count() == 1


def test_template_required_and_no_spaces():
    with pytest.raises(ValidationError):
        clean({"template": "  "})
    with pytest.raises(ValidationError):
        clean({"template": "has space"})


def test_bad_date_rejected_and_formats_accepted():
    with pytest.raises(ValidationError):
        clean({"template": "T", "created_date": "yesterday"})
    assert clean({"template": "T", "created_date": "01/09/2026"})["created_date"] == "2026-09-01"


def test_duplicate_template_warns_but_adds(store):
    store.add(ROW)
    _, warnings = store.add(ROW)
    assert warnings and store.count() == 2


def test_update_and_delete(store):
    row, _ = store.add(ROW)
    store.update(row["id"], {**ROW, "dev_name": "Dev Two"})
    assert store.matching(dev="dev two")[0]["id"] == row["id"]
    store.delete(row["id"])
    assert store.count() == 0
    with pytest.raises(KeyError):
        store.delete(row["id"])


def test_search_all_tokens_any_field(store):
    store.add(ROW)
    store.add({**ROW, "template": "TPL-B", "domain": "b.example.org", "dev_name": "Dev Two"})
    assert [r["template"] for r in store.matching(q="b.example dev two")] == ["TPL-B"]
    assert store.matching(q="nothing-here") == []


def test_blank_filter_and_blank_sorts_last(store):
    store.add(ROW)
    store.add({**ROW, "template": "TPL-B", "created_date": "", "added_to_replit": ""})
    assert [r["template"] for r in store.matching(replit="__blank__")] == ["TPL-B"]
    assert [r["template"] for r in store.matching(sort="created_date", desc=True)] == ["TPL-A", "TPL-B"]
    assert [r["template"] for r in store.matching(sort="created_date", desc=False)] == ["TPL-A", "TPL-B"]


def test_corrupt_file_does_not_crash(tmp_path):
    p = tmp_path / "bt.json"
    p.write_text("{not json", encoding="utf-8")
    assert BaseTemplateStore(p).count() == 0


def test_import_add_new_is_idempotent_and_replace_backs_up(store):
    rows = [ROW, {**ROW, "template": "TPL-B"}]
    assert store.import_rows(rows)["added"] == 2
    again = store.import_rows(rows + [{**ROW, "template": "TPL-C"}])
    assert (again["added"], again["skipped_existing"]) == (1, 2)
    res = store.import_rows([ROW], mode="replace")
    assert res["total"] == 1 and store.path.with_suffix(".bak").exists()


def test_import_reports_invalid_rows(store):
    res = store.import_rows([ROW, {"template": ""}])
    assert res["added"] == 1 and res["rejected_count"] == 1


# ── excel ────────────────────────────────────────────────────────────────────
def test_read_rows_finds_header_below_title_and_alias_headers():
    buf = make_xlsx([["2026-09-01", "TPL-A", "a.org", "Yes", "Dev", "c"], [None] * 6, ["2026-09-02", "TPL-B", None, None, None, None]], title_rows=2)
    res = excel.read_rows(buf)
    assert res["sheet"] == "Base Template" and [r["template"] for r in res["rows"]] == ["TPL-A", "TPL-B"]
    assert res["skipped_blank"] == 1 and res["rows"][0]["added_to_replit"] == "Yes"


def test_read_rows_errors_are_user_facing():
    with pytest.raises(excel.ImportError_):
        excel.read_rows(io.BytesIO(b"not an excel file"))
    wb = Workbook()
    wb.active.title = "A"
    wb.create_sheet("B")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    with pytest.raises(excel.ImportError_, match="Sheets in this file"):
        excel.read_rows(buf)
    with pytest.raises(excel.ImportError_, match="Monitoring Templates"):
        excel.read_rows(make_xlsx([], header=["x", "y"]))


def test_export_roundtrip_and_formula_guard(store):
    store.add({**ROW, "comments": "=HYPERLINK(\"http://evil\")"})
    data = excel.build_workbook(store.matching())
    ws = load_workbook(io.BytesIO(data)).active
    assert ws.title == "Base Template" and ws["B2"].value == "TPL-A"
    assert ws["F2"].value.startswith("'=")
    back = excel.read_rows(io.BytesIO(data))
    assert back["rows"][0]["comments"].startswith("=")          # guard removed again on re-import


# ── http ─────────────────────────────────────────────────────────────────────
def test_page_and_empty_list(client):
    assert b"Base Templates" in client.get("/base-templates").data
    d = client.get("/api/base-templates").get_json()
    assert d["matched"] == 0 and d["summary"]["total"] == 0


def test_changes_require_header(client):
    assert client.post("/api/base-templates", json=ROW).status_code == 400
    assert client.post("/api/base-templates", json=ROW, headers=H).status_code == 201


def test_add_then_visible_in_list_and_validation_error(client):
    client.post("/api/base-templates", json=ROW, headers=H)
    assert client.post("/api/base-templates", json={"template": ""}, headers=H).status_code == 400
    d = client.get("/api/base-templates?q=tpl-a").get_json()
    assert d["matched"] == 1 and d["rows"][0]["domain"] == "a.example.org"


def test_put_delete_404(client):
    rid = client.post("/api/base-templates", json=ROW, headers=H).get_json()["row"]["id"]
    assert client.put(f"/api/base-templates/{rid}", json={**ROW, "template": "TPL-Z"}, headers=H).status_code == 200
    assert client.delete(f"/api/base-templates/{rid}", headers=H).status_code == 200
    assert client.delete(f"/api/base-templates/{rid}", headers=H).status_code == 404
    assert client.put(f"/api/base-templates/{rid}", json=ROW, headers=H).status_code == 404


def test_bad_paging_params(client):
    assert client.get("/api/base-templates?page=abc").status_code == 400


def test_import_endpoint(client):
    buf = make_xlsx([["2026-09-01", "TPL-A", "a.org", "Yes", "Dev", ""], ["2026-09-02", "TPL-B", "b.org", "No", "Dev", ""]])
    r = client.post("/api/base-templates/import", data={"file": (buf, "x.xlsx"), "mode": "add_new"}, headers=H)
    assert r.status_code == 200 and r.get_json()["added"] == 2
    buf = make_xlsx([["2026-09-01", "TPL-A", "a.org", "Yes", "Dev", ""], ["2026-09-02", "TPL-B", "b.org", "No", "Dev", ""]])
    r2 = client.post("/api/base-templates/import", data={"file": (buf, "x.xlsx"), "mode": "add_new"}, headers=H).get_json()
    assert r2["added"] == 0 and r2["skipped_existing"] == 2
    # a row added to the sheet later shows up on the next import
    buf3 = make_xlsx([["2026-09-01", "TPL-A", "a.org", "Yes", "Dev", ""], ["2026-09-03", "TPL-NEW", "n.org", "", "", ""]])
    assert client.post("/api/base-templates/import", data={"file": (buf3, "x.xlsx")}, headers=H).get_json()["added"] == 1
    assert client.get("/api/base-templates?q=tpl-new").get_json()["matched"] == 1


def test_import_rejects_bad_input(client):
    assert client.post("/api/base-templates/import", data={}, headers=H).status_code == 400
    r = client.post("/api/base-templates/import", data={"file": (io.BytesIO(b"x"), "x.csv")}, headers=H)
    assert r.status_code == 400
    r = client.post("/api/base-templates/import", data={"file": (io.BytesIO(b"x"), "x.xlsx")}, headers=H)
    assert r.status_code == 400 and "could not be read" in r.get_json()["error"]
    empty = make_xlsx([])
    r = client.post("/api/base-templates/import", data={"file": (empty, "x.xlsx"), "mode": "replace"}, headers=H)
    assert r.status_code == 400


def test_export_endpoint_respects_filter(client):
    client.post("/api/base-templates", json=ROW, headers=H)
    client.post("/api/base-templates", json={**ROW, "template": "TPL-B"}, headers=H)
    r = client.get("/api/base-templates/export?q=tpl-b")
    ws = load_workbook(io.BytesIO(r.data)).active
    assert [c.value for c in ws["B"]][1:] == ["TPL-B"]


def test_nav_has_base_templates_in_full_app():
    import importlib
    dash = importlib.import_module("dashboard")
    c = dash.create_app().test_client()
    html = c.get("/").data.decode()
    assert 'href="/base-templates"' in html
    assert 'class="nav on" href="/base-templates"' in c.get("/base-templates").data.decode()


# ── Cached Domain format (same as the Domain column in the other tools) ──────
@pytest.mark.parametrize("raw,expected", [
    ("https://WWW.Example.GOV/path/page.html?x=1", "www.example.gov"),
    ("example.gov:8443/a", "example.gov"),
    ("www.example.gov.", "www.example.gov"),
    ("  www.example.gov  ", "www.example.gov"),
    ("https://a.org/x, B.org ;\nhttp://a.org/y", "a.org, b.org"),
    ("", ""),
])
def test_normalize_domains(raw, expected):
    assert normalize_domains(raw)[0] == expected


def test_normalize_matches_other_tabs_extract_domain():
    from url_analyzer.backlog.normalize import extract_domain
    for url in ("https://Sub.Example.org/a/b.pdf", "http://example.org:80/", "www.example.org/x"):
        assert normalize_domains(url)[0] == extract_domain(url)


def test_non_address_text_kept_and_counted():
    assert normalize_domains("N/A") == ("N/A", 1)
    assert normalize_domains("Multiple sites") == ("Multiple sites", 1)


def test_clean_and_edit_normalise_domain(store):
    assert clean({"template": "T", "domain": "HTTPS://Foo.Org/x"})["domain"] == "foo.org"
    row, _ = store.add({**ROW, "domain": "https://A.org/z"})
    assert row["domain"] == "a.org"
    assert store.update(row["id"], {**ROW, "domain": "http://B.org:81"})[0]["domain"] == "b.org"


def test_import_corrects_domains_and_dedupes_against_existing(client):
    rows = [["2026-09-01", "TPL-A", "https://WWW.A.org/path/", "Yes", "Dev", ""], ["2026-09-02", "TPL-B", "N/A", "", "", ""]]
    r = client.post("/api/base-templates/import", data={"file": (make_xlsx(rows), "x.xlsx")}, headers=H).get_json()
    assert (r["added"], r["domains_reformatted"], r["domains_not_recognised"]) == (2, 1, 1)
    assert client.get("/api/base-templates?q=tpl-a").get_json()["rows"][0]["domain"] == "www.a.org"
    # the sheet row written differently is still the same row
    again = [["2026-09-01", "TPL-A", "www.a.org", "Yes", "Dev", ""]]
    assert client.post("/api/base-templates/import", data={"file": (make_xlsx(again), "x.xlsx")}, headers=H).get_json()["added"] == 0


def test_old_stored_rows_are_normalised_on_load(tmp_path):
    import json
    p = tmp_path / "bt.json"
    p.write_text(json.dumps({"rows": [{"id": "1", "seq": 1, "template": "T", "domain": "HTTPS://X.org/a", "created_date": "",
                                        "added_to_replit": "", "dev_name": "", "comments": ""}]}), encoding="utf-8")
    assert BaseTemplateStore(p).matching()[0]["domain"] == "x.org"
